"""Build a numerical audit and matching Markdown/Word runtime explanation."""
from datetime import datetime
from zoneinfo import ZoneInfo
import re
import statistics
from common import *
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn


def main():
    home=WORK/'amplify_runtime_audit'
    probes={t:load(home/f'probe_threads{t}.json') for t in [2,4]}
    observed=[]
    for path in sorted((WORK/'external_runs/AMPlify').glob('seed*/training.log')):
        raw=path.read_text(errors='replace')
        parsed=re.findall(r'363/363 - (\d+)s - loss: ([.\d]+) - val_loss: ([.\d]+)',raw)
        rows=[{'epoch':i+1,'seconds':int(sec),'loss_rounded':float(loss),'val_loss_rounded':float(val)}
              for i,(sec,loss,val) in enumerate(parsed)]
        if rows:
            (home/(path.parent.name+'_training_log_snapshot.txt')).write_text(raw)
            observed.append({'seed':int(path.parent.name[4:]),'completed_epochs':len(rows),
                'epoch_rows':rows,'log_sha256':sha(path)})
    stamp=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(timespec='seconds')
    dump(home/'audit_snapshot.json',{'captured_at':stamp,'formal_training_logs':observed,
        'probe_files':[f'probe_threads{t}.json' for t in probes],
        'probe_scope':'CPU only; archived checkpoint; tiny diagnostic updates never saved to formal model files',
        'drivers_differ_only_by_negative_count_7962_to_7949':all(
            (WORK/f'external_runs/AMPlify/seed{s}/workspace/baseline/sota_bench/retrain_amplify_regr_neg.py').read_text()
            == (ROOT/'baseline/sota_bench/retrain_amplify_regr_neg.py').read_text().replace('len(neg) == 7962','len(neg) == 7949')
            for s in [120,121,123])})
    blocks=[]
    def h(v):blocks.append(('h',v))
    def p(v):blocks.append(('p',v))
    def t(v):blocks.append(('t',v))
    h('AMPlify训练耗时与代码逻辑核查')
    p('核查时间：'+stamp+'。对象为当前v15含负样本重训；诊断没有修改正式训练权重、学习率、batch size、早停设置或训练数据。')
    h('结论')
    p('目前未发现重复嵌套训练、数据重复加载造成额外整轮训练、早停未连接或梯度不更新等能够解释异常长耗时的循环错误。观察到的耗时主要由CPU执行、全参数BiLSTM/注意力网络、训练样本和批次数、长patience共同解释。传统深度学习方法仍可能有较大的训练计算量，208条肽的单次前向推理时间不能直接预测完整训练时间。')
    p('同时确认三项需明确记录的实现细节：早停恢复不保证绝对最小验证损失；旧计时脚本没有在代码内锁定TensorFlow线程数；自定义注意力层复用同一initializer，在当前Keras下会产生相同初始投影矩阵。这些细节与训练循环故障分开判断，不能据此直接归因长耗时。')
    h('网络规模与工作量')
    t([['项目','实际配置'],['网络','Masking → BiLSTM（每方向512单元）→ 32头注意力 → context attention → sigmoid回归头'],
       ['总可训练参数',f"{probes[2]['parameters']:,}"],
       ['BiLSTM参数','2,183,168'],['多头注意力参数','4,198,400'],['context attention参数','1,050,624'],['输出层参数','1,025'],
       ['输入长度／batch size','50／32'],['含负样本语料','4,929条正序列＋7,949条负序列＝12,878条'],
       ['每种子划分','11,591条训练＋1,287条验证'],['每轮批次','363个训练批次＋41个验证批次'],
       ['外部推理','208条肽，默认batch 32时为7个前向批次'],['当前计算环境','TensorFlow 2.12.1，CPU；主要计算线程2，inter-op线程1'],
       ['宿主机','AMD Ryzen 9 9950X；16物理核、32逻辑线程']])
    p('模型结构与本地作者仓库train_amplify.py一致；回归损失、标签缩放、单模型/种子和长度50属于既有适配。Masking保证掩码参与计算语义，并不意味着所有padding矩阵运算都被跳过。每轮训练除前向传播，还包含反向传播、Adam更新和完整验证。')
    h('实测运行情况')
    rows=[['种子','已完成轮数','已完成轮耗时范围','最近三轮耗时（秒）']]
    for item in observed:
        secs=[r['seconds'] for r in item['epoch_rows']]
        rows.append([str(item['seed']),str(item['completed_epochs']),f'{min(secs)}–{max(secs)}秒',', '.join(map(str,secs[-3:]))])
    t(rows)
    p('早期两个AMPlify任务并行时每轮约132–154秒。CPU并发扩展、第三个AMPlify加入及本次短诊断期间，近期部分轮增至约210–225秒，说明单任务速度对资源竞争敏感；这些观测不能把变化全部归因于某一个因素。32逻辑线程不是32个独立物理核，不能按任务数直接线性推算吞吐。')
    h('隔离诊断：训练和前向均正常执行')
    rows=[['诊断线程数','208肽前向中位数','训练batch中位数','验证batch中位数','按批次外推单轮（非完整实测）']]
    for threads,d in probes.items():
        rows.append([str(threads),f"{statistics.median(d['inference_208_seconds']):.3f}秒",
            f"{statistics.median(d['train_batch_seconds']):.3f}秒",f"{statistics.median(d['val_batch_seconds']):.3f}秒",
            f"{d['estimated_epoch_seconds']:.1f}秒"])
    t(rows)
    improvement=1-statistics.median(probes[4]['train_batch_seconds'])/statistics.median(probes[2]['train_batch_seconds'])
    p(f'使用归档seed123权重，分别在独立2线程和4线程进程执行208肽前向及256训练样本/64验证样本的3轮诊断；第一轮编译预热不计batch统计。两个诊断均损失有限、输出层权重实际更新。4线程的训练batch中位数约下降{improvement:.1%}，说明2线程预算确实限制速度；由于诊断顺序执行且正式作业持续运行，这只是当前负载下的小样本观察，不能作为最终论文计时或保证持续提速比例。')
    p('旧归档外测推理中位数为0.1859秒，脚本同样重建AMPlify并加载seed123权重，使用CPU执行；但计时代码没有显式固定TensorFlow intra/inter-op线程数，历史进程的实际线程预算未在该结果中完整记录。当前并发环境下的2线程诊断约0.9秒，说明不同资源条件下的0.1859秒不能直接代入训练时间估算，也不能据此判定旧测量错误。v15最终计时仍按已批准方案在设备空闲时统一线程预算重测。')
    h('早停核查')
    p('当前配置为max_epochs=1000、monitor=val_loss、patience=50、min_delta=0.001、restore_best_weights=True。回调已传入model.fit，日志确有val_loss；单任务命令只包含一个种子，没有意外把10种子或作者5模型集成重复嵌套。1000是上限，不是计划必跑轮数。patience=50要求连续50轮没有达到阈值的改善，因此即使很早稳定，也需至少约51轮。以60轮为例，每轮140秒对应2.33小时，每轮210秒对应3.50小时。')
    p('当前Keras 2.12.1用同一min_delta判据更新best_weights：只有验证损失相对已认可最佳值下降超过0.001，才更新恢复点。因此“restore_best_weights=True”不等于恢复所有轮次中绝对最小val_loss。用安装的真实回调输入预设损失序列实证：第2轮0.0048被认可，第5轮降到0.0039仍不足0.001改善；回调在第52轮停止并恢复第2轮。该行为符合现有代码和登记的阈值，应在方法学中精确表述。')
    p('标签为y=(log10MIC+4)/9，所以0.001的缩放MSE变化对应0.081的log10MIC域MSE变化。它不是“0.1%相对改善”，也不能与原作者分类准确率的0.001直接作数值含义上的等同。改变patience或min_delta会改变已登记训练配方，本次诊断没有作此修改。运行日志只输出四位小数，不能用日志边界数值精确重建每一轮早停内部状态。')
    h('注意力初始化兼容性细节')
    p('本地作者layers.py在MultiHeadAttention.build内把同一个GlorotNormal实例用于Wq、Wk、Wv、Wo。当前Keras明确警告该实例多次调用会重复取值；两次独立诊断都证实四个1024×1024初始矩阵逐位相同。它们是独立变量，不是训练过程中共享同一权重，不能据初值相同推断梯度不更新或模型无法学习。')
    p('这是原有实现与当前框架组合下的已确认初始化行为，需要记录并单独评估方法忠实性；目前没有因果证据表明它导致本轮耗时长。本轮修复生成的三个运行驱动与原驱动相比，仅把负样本数断言7962改为7949，未引入该初始化行为。若之后决定改为独立初始化器，将属于另一次训练实现修订，应统一影响条件和种子，不能只修改当前尚未开始的部分种子。')
    h('处理建议')
    p('继续保留正在正常推进的正式任务。优先优化执行资源：为AMPlify设置单独的重型CPU任务并发预算，并以总吞吐而非单任务CPU占用评估2线程/4线程组合。小诊断尚不足以证明四个重型任务同时运行优于两个。')
    p('若采用GPU加速，应使用独立兼容的TensorFlow GPU环境先验证前向、损失、更新、掩码和显存预算，再安排完整种子训练；不能直接改写正在使用的环境。缩小BiLSTM、增大batch或缩短patience会改变实验配方，不作为本次运行优化的直接手段。')
    h('证据文件')
    t([['内容','相对项目根目录的路径'],['原训练驱动','baseline/sota_bench/retrain_amplify_regr_neg.py'],['作者训练与层实现','baseline/sota_bench/amplify_repo/src/train_amplify.py；src/layers.py'],['正式训练日志','artifacts/repair_v15/external_runs/AMPlify/seed*/training.log'],['诊断代码','scripts/repair_v15/diagnose_amplify_runtime.py'],['2/4线程诊断结果','artifacts/repair_v15/amplify_runtime_audit/probe_threads2.json；probe_threads4.json'],['本次日志快照与审计结果','artifacts/repair_v15/amplify_runtime_audit/audit_snapshot.json'],['原推理计时实现','scripts/eval_cost_timing_new4_tf.py']])
    doc=Document(); normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    md=[]; expected=[]
    for kind,value in blocks:
        if kind=='h':doc.add_heading(value,level=1);md.extend(['# '+value,'']);expected.append(value)
        elif kind=='p':doc.add_paragraph(value);md.extend([value,'']);expected.append(value)
        else:
            table=doc.add_table(rows=0,cols=len(value[0]));table.style='Table Grid'
            for row in value:
                for cell,text in zip(table.add_row().cells,row):cell.text=text;expected.append(text)
            md.extend(['| '+' | '.join(value[0])+' |','| '+' | '.join(['---']*len(value[0]))+' |'])
            md.extend('| '+' | '.join(row)+' |' for row in value[1:]);md.append('')
    base=REPORT/'AMPlify_runtime_review_ZH';base.with_suffix('.md').write_text('\n'.join(md));doc.save(base.with_suffix('.docx'))
    check=Document(base.with_suffix('.docx'))
    actual='\n'.join([p.text for p in check.paragraphs]+[c.text for t in check.tables for row in t.rows for c in row.cells])
    assert all(x in actual for x in expected)
    print(base)


if __name__=='__main__':main()
