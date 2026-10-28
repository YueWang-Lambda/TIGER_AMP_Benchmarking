"""Verify the live CPU-to-GPU handoff and emit matching Markdown/Word records."""
from datetime import datetime
from zoneinfo import ZoneInfo
import csv
import re
import subprocess
from common import *
from external_one import tf_gpu_profile
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn


def main():
    current=load(WORK/'tf_gpu_migration/current.json');directory=Path(current['directory'])
    profile=tf_gpu_profile();state=load(WORK/'queue/campaign/state.json')
    assert state['pid']==current['new_controller_pid']
    assert state['status']=='RUNNING' and not state['failures']
    jobs={j['id']:j for j in load(WORK/'campaign_jobs.json')}
    affected={k:v for k,v in jobs.items() if v.get('method') in ['AMPlify','MBC-Attention']}
    assert len(affected)==20 and all(j['resource']=='gpu' for j in affected.values())
    assert not any(a['resource']=='cpu' and a['id'] in affected for a in state['active'])
    preserved=load(directory/'completed_preserved_hashes.json')
    for folder,files in preserved.items():
        for name,h in files.items():assert sha(ROOT/folder/name)==h,(folder,name)
    cancelled=load(directory/'cancelled_workers.json')
    for entry in cancelled:
        for pid in entry['descendant_pids']:
            p=Path(f'/proc/{pid}/stat')
            assert not p.exists() or p.read_text().rsplit(')',1)[1].split()[0]=='Z',pid
    handoff=[]
    cancelled_ids={e['id'] for e in cancelled}
    for old in load(directory/'workers_before.json'):
        if old['id'] in cancelled_ids:continue
        if Path(jobs[old['id']]['completion']).exists():
            handoff.append({'id':old['id'],'status':'COMPLETED'});continue
        fields=Path(f"/proc/{old['pid']}/stat").read_text().rsplit(')',1)[1].split()
        assert fields[0]!='Z' and fields[19]==old['start_ticks']
        handoff.append({'id':old['id'],'pid':old['pid'],'status':'ORIGINAL_PROCESS_PRESERVED'})
    running={a['id']:a for a in state['active']};methods={};rows=[]
    for method in ['AMPlify','MBC-Attention']:
        cpu=[];gpu_done=[];gpu_active=[];pending=[]
        for seed in range(120,130):
            job=WORK/f'external_runs/{method}/seed{seed}';ident=f'{method}_with_negatives_seed{seed}'
            if str(job.relative_to(ROOT)) in preserved:
                cpu.append(seed);continue
            if (job/'DONE.json').exists():
                done=load(job/'DONE.json');assert done['status']=='PASS' and done['training_device'].startswith('GPU:')
                assert load(job/'metrics.json')['n_pairs']==2211
                gpu_done.append(seed)
            elif ident in running:
                gpu_active.append(seed)
            else:pending.append(seed)
            if (job/'runtime_environment.json').exists():
                runtime=load(job/'runtime_environment.json')
                assert runtime['execution']==profile['execution'][method] and runtime['memory_cap_mib']==1024
                assert runtime['prefix'].endswith('/mbc_att_gpu_v15')
                assert load(job/'invocation.json')['env']['CUDA_VISIBLE_DEVICES'] in ['0','1']
                log=job/'training.log';lines=log.read_text(errors='replace').splitlines() if log.exists() else []
                epochs=[v for v in lines if re.match(r'^\d+/\d+ - ',v)]
                rows.append([method,str(seed),'已完成训练与2,211边外测' if seed in gpu_done else 'GPU训练中',
                    str(len(epochs)),epochs[-1] if epochs else '首轮尚未结束'])
        methods[method]={'preserved_completed_cpu_seeds':cpu,'completed_gpu_seeds':gpu_done,
            'active_gpu_seeds':gpu_active,'pending_gpu_seeds':pending}
    query=subprocess.check_output(['nvidia-smi','--query-gpu=index,uuid,memory.used,memory.total',
        '--format=csv,noheader,nounits'],text=True)
    cards={r[1].strip():{'index':int(r[0]),'used_mib':int(r[2]),'total_mib':int(r[3]),'processes':[]}
           for r in csv.reader(query.splitlines())}
    query=subprocess.check_output(['nvidia-smi','--query-compute-apps=gpu_uuid,pid,used_memory',
        '--format=csv,noheader,nounits'],text=True)
    for r in csv.reader(query.splitlines()):cards[r[0].strip()]['processes'].append({'pid':int(r[1]),'memory_mib':int(r[2])})
    assert all(len(c['processes'])<=5 for c in cards.values())
    stamp=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(timespec='seconds')
    snapshot={'status':'PASS','captured_at':stamp,'controller_pid':state['pid'],
        'cancelled_cpu_attempts':len(cancelled),'methods':methods,'unaffected_workers':handoff,
        'completed_cpu_files_unchanged':True,'no_tf_cpu_jobs':True,'gpu_cards':cards,
        'training_observations':rows,'queue_state':state}
    dump(directory/'handoff_verification.json',snapshot)
    current.update(status='GPU_MIGRATION_VERIFIED',verification=str(directory/'handoff_verification.json'))
    dump(WORK/'tf_gpu_migration/current.json',current)
    ready=load(WORK/'tf_gpu_repair/RUNTIME_READY.json')
    ready.update(status='ACTIVE_GPU_QUEUE',tensorflow_allocation_cap_mib=1024,
        proposed_full_process_reservation_mib=2048,active_profile=str(WORK/'tf_gpu_runtime.json'),
        migration_verification=str(directory/'handoff_verification.json'))
    dump(WORK/'tf_gpu_repair/RUNTIME_READY.json',ready)
    blocks=[]
    def h(s):blocks.append(('h',s))
    def p(s):blocks.append(('p',s))
    def t(r):blocks.append(('t',r))
    h('AMPlify与MBC-Attention正式训练迁移至GPU')
    p('核验时间：'+stamp+'。依据用户“取消AMPlify和MBC-Attention的CPU训练任务，切换为GPU训练”的指令，迁移已完成并通过进程、运行环境、队列与文件保全核验。')
    h('处理范围')
    p('停止并归档3个尚未完成的AMPlify CPU尝试：seed123、seed122、seed124。MBC-Attention seed125在切换前已正常完成，计入保留结果。保留AMPlify seed120/121共2个，以及MBC-Attention seed120–125共6个完整CPU结果；已完成文件逐字节未变。余下AMPlify 8个、MBC-Attention 4个种子均迁入GPU队列，不再调度这两种方法的CPU训练。')
    p('被中止的原驱动没有完整保存可恢复的优化器、早停与随机状态，因此GPU任务按原登记种子重新初始化训练；不是从CPU中止轮次无缝续训。种子集合、训练/验证划分算法、样本、模型结构、损失、学习率、batch size、最大轮数和早停设置均保留。中止日志及工作目录移入本次迁移归档，写入CANCELLED.json，不混入正式结果。')
    t([['方法','保留已完成CPU种子','已完成GPU种子','GPU运行中','GPU待启动']]+[
        [m,', '.join(map(str,r['preserved_completed_cpu_seeds'])),', '.join(map(str,r['completed_gpu_seeds'])) or '—',
         ', '.join(map(str,r['active_gpu_seeds'])) or '—',', '.join(map(str,r['pending_gpu_seeds'])) or '—'] for m,r in methods.items()])
    h('实际运行配置')
    t([['项目','配置'],['环境','/home/ubuntu/anaconda3/envs/mbc_att_gpu_v15'],
       ['AMPlify执行方式','GPU eager；避开已实证的带掩码反向LSTM图模式差异，保留cuDNN'],
       ['MBC-Attention执行方式','GPU graph；迁移前预测、梯度及训练更新复验通过'],
       ['精度','FP32；TF32关闭'],['线程','intra-op 2；inter-op 1'],
       ['TensorFlow分配上限','每任务1,024 MiB'],['整进程调度预留','每任务2,048 MiB，另保留每卡整体安全余量'],
       ['GPU并发','每卡最多5个计算进程；显存预算与进程数同时限制'],
       ['CPU方法','EIPpred、AMPActiPred继续并行，占用释放出的CPU名额']])
    p('迁移前在1 GiB分配上限下分别重跑真实模型验收：208肽与2,211边预测核对通过，全部参数的确定性梯度核对通过，实际3轮训练损失有限且权重更新。AMPlify张量峰值约375 MiB，MBC-Attention约122 MiB；CUDA上下文另占显存，因此调度按整进程2 GiB预留，而不是只按张量峰值计算。')
    p('新任务通过专用启动器显式设置CUDA库路径、环境前缀、显存上限和执行方式，GPU不可用时直接报错；每个种子生成runtime_environment.json。完成记录补写实际训练GPU、环境、执行方式及运行环境文件哈希。已完成CPU种子与新GPU种子均按实际设备登记，不将整组描述为全部GPU训练。')
    h('切换与运行核验')
    p(f'原调度器先停止新增任务，再由新调度器PID {state["pid"]}接管。仅向指定AMPlify任务进程组发出终止信号；{len(handoff)}个其他活跃工作进程保留原PID继续运行，或在交接后正常完成。队列当前{state["status"]}，失败记录{len(state["failures"])}项；全部20个TensorFlow任务定义均标记GPU，已完成CPU任务通过完成记录跳过，未重复训练。')
    t([['GPU','计算进程数','当前已用显存（MiB）']]+[[str(c['index']),str(len(c['processes'])),str(c['used_mib'])] for c in cards.values()])
    t([['方法','种子','状态','已完成epoch数','最近完成轮的日志']]+rows)
    p('以上速度来自共享设备上的当前训练观察，不能直接当作论文推理时间；最终推理仍按统一设备、线程及完整外部终点协议，在设备空闲时重测。GPU训练的浮点顺序和随机实现可能影响具体学习轨迹，故保留逐种子实际运行环境，不承诺CPU与GPU训练权重逐位一致。')
    h('交付与续作位置')
    t([['内容','相对项目根目录的路径'],['迁移核验与归档',str(directory.relative_to(ROOT))],
        ['GPU运行配置','artifacts/repair_v15/tf_gpu_runtime.json'],['统一队列','artifacts/repair_v15/queue/campaign/state.json'],
        ['各种子结果与运行环境','artifacts/repair_v15/external_runs/<方法>/seed<种子>/'],
        ['GPU数值验收','scripts/repair_v15/tf_gpu_migration_preflight.py'],
        ['迁移核验与报告生成','scripts/repair_v15/tf_gpu_migration_report.py']])
    doc=Document();style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    md=[];expected=[]
    for kind,value in blocks:
        if kind=='h':doc.add_heading(value,1);md.extend(['# '+value,'']);expected.append(value)
        elif kind=='p':doc.add_paragraph(value);md.extend([value,'']);expected.append(value)
        else:
            table=doc.add_table(rows=0,cols=len(value[0]));table.style='Table Grid'
            for row in value:
                for cell,s in zip(table.add_row().cells,row):cell.text=s;expected.append(s)
            md.extend(['| '+' | '.join(value[0])+' |','| '+' | '.join(['---']*len(value[0]))+' |'])
            md.extend('| '+' | '.join(row)+' |' for row in value[1:]);md.append('')
    base=REPORT/'TensorFlow_GPU_migration_ZH';base.with_suffix('.md').write_text('\n'.join(md));doc.save(base.with_suffix('.docx'))
    check=Document(base.with_suffix('.docx'));actual='\n'.join([p.text for p in check.paragraphs]+[c.text for t in check.tables for row in t.rows for c in row.cells])
    assert all(s in actual for s in expected)
    print({'status':'PASS','methods':methods,'report':str(base)})


if __name__=='__main__':main()
