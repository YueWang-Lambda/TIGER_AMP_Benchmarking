"""Read-only audit of terminal annotations in the active v15 training inputs."""
import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from common import ROOT, REPORT, sha, dump

import numpy as np
import pandas as pd
import torch
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

from code.dataloader import preprocess_dataset
from code.utils.config import load_config
from code.utils.features import calculate_property, terminal_charge_adjustment


def main():
    torch.set_num_threads(1)
    out = REPORT / 'terminal_modifications_20261012'
    out.mkdir(parents=True, exist_ok=True)
    rawdir = ROOT.parents[2] / 'POAP/Data/dbaasp_jsons'
    upstream = rawdir.parent / 'dbaasp_amp_training_by_cfu_group_ug_per_mL.csv'
    historical = ROOT.parents[1] / 'newdata/dbaasp_amp_training_by_cfu_group_ug_per_mL _20260714.csv'
    source = ROOT / 'metadata/train_val_by_cfu_group_ug_per_mL.csv'
    external = ROOT / 'reports/methods_package_S4_v1/4_S4_benchmark/data/train_restricted4929_tiger.csv'
    config = ROOT / 'runs_repair_v15/S4/S4_seed120/config.json'
    feature = ROOT / 'code/utils/features.py'
    loader = ROOT / 'code/dataloader.py'
    inputs = {str(p): sha(p) for p in [source, upstream, historical, external, config, feature, loader]}
    src = pd.read_csv(source)
    ids = set(src.dbaasp_id)
    rawrows = []
    for path in sorted(rawdir.glob('*.json')):
        blob = path.read_bytes()
        j = json.loads(blob)
        n, c = j.get('nTerminus') or {}, j.get('cTerminus') or {}
        rawrows.append(dict(dbaasp_id=j.get('dbaaspId'), sequence=j.get('sequence'),
                            n_terminus=n.get('name', ''), c_terminus=c.get('name', ''),
                            n_description=n.get('description', ''), c_description=c.get('description', ''),
                            n_is_null=j.get('nTerminus') is None, c_is_null=j.get('cTerminus') is None,
                            raw_json_path=str(path)))
        if j.get('dbaaspId') in ids:
            import hashlib
            inputs[str(path)] = hashlib.sha256(blob).hexdigest()
    raw = pd.DataFrame(rawrows)
    assert not raw.dbaasp_id.duplicated().any()
    joined = src.merge(raw, on='dbaasp_id', how='left', suffixes=('', '_raw'), validate='many_to_one', indicator=True)
    assert joined._merge.eq('both').all()
    assert joined.sequence.str.strip().eq(joined.sequence_raw.str.strip()).all()
    mismatches = joined.loc[
        joined.n_terminus.fillna('').ne(joined.n_terminus_raw.fillna('')) |
        joined.c_terminus.fillna('').ne(joined.c_terminus_raw.fillna(''))
    ].copy()
    assert set(mismatches.dbaasp_id) == {'DBAASPS_6804', 'DBAASPS_6808'}
    mismatches.drop(columns=['_merge']).to_csv(out / 'two_AM_to_AMD_annotation_changes.csv',index=False,encoding='utf-8-sig')
    trace = raw[raw.dbaasp_id.isin(ids)]
    trace.to_csv(out / 'training_10929_raw_terminal_annotations.csv', index=False, encoding='utf-8-sig')

    cfg = load_config(config)
    effective, stats = preprocess_dataset(str(source), str(ROOT / cfg.train_pdb_dir), cfg, 'train')
    cachepath = next((ROOT / 'artifacts/repair_v15/graph_cache').glob('training*.pt'))
    inputs[str(cachepath)] = sha(cachepath)
    cache = torch.load(cachepath, map_location='cpu', weights_only=False)
    cols = ['sequence', 'n_terminus', 'c_terminus', 'cfu_group', 'MIC_Escherichia_coli']
    assert cache['rows'][cols].fillna('').equals(effective[cols].fillna(''))
    assert len(cache['graphs']) == len(effective) == 7217
    for r in mismatches.itertuples():
        matching = effective[effective.sequence.eq(r.sequence) & effective.cfu_group.eq(r.cfu_group)]
        assert len(matching) == 1
        assert matching.c_terminus.iloc[0] == r.c_terminus
        assert matching.MIC_Escherichia_coli.iloc[0] == r.MIC_Escherichia_coli
    effective[cols].to_csv(out / 'TIGER_7217_effective_terminal_annotations.csv', index=False, encoding='utf-8-sig')

    frames = {
        'raw_JSON_unique_peptides': raw,
        'upstream_full_CFU_rows': pd.read_csv(upstream),
        'historical_selected_CFU_rows': pd.read_csv(historical, encoding='gb18030'),
        'current_TIGER_source_rows': src,
        'current_TIGER_effective_observations': effective,
        'external_benchmark_positive_source_rows': pd.read_csv(external),
        'historical_non_CFU_table_not_used_by_this_config': pd.read_csv(ROOT / 'metadata/train_val_ug_per_mL.csv'),
    }
    totals, inventory = {}, []
    for label, frame in frames.items():
        n, c = frame.n_terminus.fillna(''), frame.c_terminus.fillna('')
        other = (~n.isin(['', 'ACT'])) | (~c.isin(['', 'AMD']))
        totals[label] = dict(rows=len(frame), unique_sequences=int(frame.sequence.nunique()),
                             n_ACT=int(n.eq('ACT').sum()), c_AMD=int(c.eq('AMD').sum()),
                             n_missing=int(n.eq('').sum()), c_missing=int(c.eq('').sum()),
                             other_terminal_rows=int(other.sum()),
                             n_nonempty_categories=int(n[n.ne('')].nunique()),
                             c_nonempty_categories=int(c[c.ne('')].nunique()))
        for side, values in [('N', n), ('C', c)]:
            for term, count in values.value_counts().items():
                inventory.append(dict(dataset=label, terminal=side, annotation=term or '<missing>', count=int(count)))
    pd.DataFrame(inventory).to_csv(out / 'terminal_annotation_inventory.csv', index=False, encoding='utf-8-sig')
    joint = effective.fillna('').groupby(['n_terminus', 'c_terminus']).size().reset_index(name='observations')
    joint.to_csv(out / 'TIGER_terminal_joint_counts.csv', index=False, encoding='utf-8-sig')
    assert totals['current_TIGER_source_rows']['other_terminal_rows'] == 0
    assert totals['external_benchmark_positive_source_rows']['other_terminal_rows'] == 0

    seq = 'RFWWFRRR'
    cases = [('', ''), ('ACT', 'AMD'), ('ACT', ''), ('acetyl', ''), ('FOR', ''), ('formyl', ''),
             ('Pyr', ''), ('pyroglutamate', ''), ('C8', ''), ('PEG 5000', ''),
             ('', 'AMD'), ('', 'AM'), ('', 'OMe'), ('', 'methyl ester'), ('', 'OBzl'), ('', 'C12H25-NH2')]
    probes = []
    base = np.array(calculate_property(seq, '', '', scaling='raw'))
    for n, c in cases:
        p = np.array(calculate_property(seq, n, c, scaling='raw'))
        ix = np.flatnonzero(p != base).tolist()
        assert set(ix).issubset({7, 9})
        probes.append(dict(sequence=seq, n_input=n, c_input=c,
                           code_adjustment=terminal_charge_adjustment(n, c),
                           charge=p[7], charge_density=p[9], changed_property_indices=str(ix)))
    pd.DataFrame(probes).to_csv(out / 'terminal_code_behavior_probes.csv', index=False, encoding='utf-8-sig')
    # These are software-behaviour probes, not assertions of chemical validity.
    assert terminal_charge_adjustment('FOR', '') == terminal_charge_adjustment('', '')
    assert terminal_charge_adjustment('', 'OMe') == terminal_charge_adjustment('', '')
    assert terminal_charge_adjustment('ACT', '') == terminal_charge_adjustment('acetyl', '')
    assert calculate_property(seq,'ACT','AMD') == calculate_property(seq,'','')

    blocks = []
    def h(t): blocks.append(('h', t))
    def p(t): blocks.append(('p', t))
    def table(rows): blocks.append(('t', [[str(x) for x in row] for row in rows]))
    h('N、C端化学修饰：当前训练输入与编码核查')
    p('核查时间：' + datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(timespec='seconds') + '。范围为本次v15任务当前配置读取的训练数据、图缓存及其原始DBAASP JSON；只读检查，不修改训练源表、模型或任务队列。')
    p('结论：当前CSV的末端字段只有ACT、AMD和空值；但原始JSON逐条核查发现两条AM被改写为AMD，且两条确实进入TIGER有效训练池，所以不能据CSV声称所有原始末端标注只有乙酰化/酰胺化。其他大多数已标注修饰没有进入当前队列。代码没有实现通用末端化学表示；少数名称仅影响电荷近似，未识别名称静默走默认分支。空值只能报告为“未标注”，不能推断为实验已证实游离端。')
    h('一、实际进入当前训练流程的末端标注')
    table([['数据层','观察/记录数','N端ACT','C端AMD','CSV中其他末端标注']] +
          [[label, totals[key]['rows'], totals[key]['n_ACT'], totals[key]['c_AMD'], totals[key]['other_terminal_rows']] for label,key in [
              ('当前TIGER源表','current_TIGER_source_rows'), ('当前TIGER有效训练/验证池','current_TIGER_effective_observations'),
              ('外部基准共同正样本源表','external_benchmark_positive_source_rows')]])
    p('ACT和AMD可以在同一条肽上同时出现，两列不能相加作为肽总数。当前7,217是按sequence×CFU整理后的训练/验证观察数，不是全部独立化学分子的数量；跨验证折分配。有效池与正在使用的已修复ACT特征图缓存逐行一致。')
    table([['N端字段','C端字段','有效观察数']] + [[r.n_terminus or '未标注', r.c_terminus or '未标注', r.observations] for r in joint.itertuples()])
    p(f'当前11,516条源记录对应{len(ids):,}个DBAASP条目；全部找到原始JSON。N端字段全部一致；C端有两条AM→AMD差异，其余一致；序列去除首尾空白后全部一致。没有发现其他末端修饰被改成空值，但不能排除原始数据库自身漏注或仅在论文中描述修饰。')
    table([['DBAASP条目','名称','序列','原始C端','训练表C端','MIC（µg/mL）'],
           ['DBAASPS_6804','Adepantin-2 (AM)','GIGKHVGKALKGLKGLLKGLGEC','AM','AMD','2.276812'],
           ['DBAASPS_6808','Adepantin-3 (AM)','GLKGLLGKALKGIGKHIGKAQGC','AM','AMD','8.991095']])
    p('两个原始JSON均将AM描述为“Acetamide; CH3CONH2”，而AMD的描述为“Amide, NH2”。AM→AMD的改写已存在于历史筛选表，非本次ACT修复造成。两条在当前TIGER有效池中均以AMD编码；外部基准正样本表还保留DBAASPS_6804。不能仅凭名称判定这属于合理别名归并还是化学身份错误，应核对对应原始实验文献后裁定；本报告不擅自改回AM或据此启动新重训。')
    h('二、其他修饰在哪个数据阶段消失')
    a=totals['upstream_full_CFU_rows']; b=totals['historical_selected_CFU_rows']; old=totals['historical_non_CFU_table_not_used_by_this_config']
    p(f'上游完整CFU表有{a["rows"]:,}行，包含{a["other_terminal_rows"]:,}行ACT/AMD以外的末端标注，如C8、C16、FOR、PEG 5000、OMe、OBzl。历史筛选表“dbaasp_amp_training_by_cfu_group_ug_per_mL _20260714.csv”有{b["rows"]:,}行，其末端字段已只剩ACT、AMD和空值。除了上述两条被改写的AM记录，原始JSON对照未发现其他非ACT/AMD末端类型进入当前源表；不能将此归功于当前训练加载器自动识别并排除了它们。')
    p('已检查的上游JSON提取脚本只在read_monomers中筛选Monomer，并直接保留nTerminus/cTerminus名称；当前dataloader也没有末端白名单。此次没有定位到生成历史筛选表的完整终端筛选脚本/日志，所以只能确认文件间的实际变化，不把具体筛选原因写成已证实事实。')
    p(f'仓库另有旧表metadata/train_val_ug_per_mL.csv（{old["rows"]:,}行），仍含{old["other_terminal_rows"]:,}行其他末端修饰标注。本次配置实际读取的是train_val_by_cfu_group_ug_per_mL.csv，不是这张旧表；不能据旧表断定这些肽参加了本次训练，也不能把本结论推广为所有历史实验均未使用它们。')
    h('三、模型如何处理修饰输入')
    table([['通道/步骤','当前实现及边界'],
           ['序列通道','20种标准氨基酸token；没有末端修饰token。相同序列的不同末端修饰不会得到不同序列编码。'],
           ['结构通道','按序列查找PDB，只把具有Cα的标准氨基酸建为残基节点；没有为末端附加化学基团建立原子图或专属节点。同序列/不同末端使用同名PDB。'],
           ['理化通道','修饰字段只改变十个理化特征中的charge与charge_density；pI、GRAVY等仍由标准序列计算。没有通用质量、脂肪链、PEG或其他基团表示。'],
           ['N端名称识别','ACT先映射为acetyl；随后检查ace/acetyl/pyro/pglu/formyl子串。命中时不添加默认的N端+1项。'],
           ['C端名称识别','检查amd/amid/amide/nh2/methyl/ester子串。命中时不添加默认的C端−1项。'],
           ['未识别及空值','没有报错或主动拒绝；走与默认端相同的电荷分支。上述处理是代码中的整数电荷近似，不是完整修饰分子的化学计算。']])
    p('实测例子：N端FOR与formyl在当前函数中的结果不同，前者不会命中；C端OMe与methyl ester也不同，前者不会命中。C8、PEG 5000、OBzl同样没有独立表示；含NH2的复杂C端名称可能被统一归入抑制默认−1项的分支。探针CSV仅展示软件行为，不代表这些字符串或整数修正已构成合理化学模型。')
    p('表示能力还有一个直接限制：同一序列的ACT+AMD与两端空值在当前函数中总电荷修正都为0，十个理化特征完全相同。这说明识别了名称也不等于模型能区分所有修饰组合；相同序列、相同CFU和同名PDB下，没有其他独立末端特征补足这一差别。')
    p('当前去重键为sequence×CFU，不含末端修饰；同一序列、同一CFU下的不同末端化学实体可在选择最小MIC代表时合并。该事实与前次ACT质量换算审计发现的两组代表切换一致，应与标签修复一起明确数据身份定义。')
    h('四、对本次修复与方法学措辞的含义')
    p('本次新增两条AM→AMD的原始标注差异，需先通过原始实验文献判断，不应立即断言需要再次全量重训。若当前AMD正确，可补充有依据的更正记录；若分子身份或标签确需修改，则把修复并入正在讨论的数据冻结，按真实输入/标签依赖核定任务。已发现的ACT分子量换算问题仍成立。建议数据入口区分已支持、未支持和未标注末端，未支持类型记录并拒绝或另行建模，避免静默当作默认端。')
    p('在解决两条原始标注差异后，方法学应准确披露当前数据终端范围与规范化依据。当前实现中，末端字段通过启发式电荷及电荷密度项影响理化输入，序列和结构通道未显式编码末端附加基团；不要表述为“模型支持所有N/C端化学修饰”。负样本的空末端字段是加载器设定，不能证明其来源分子的实验末端化学状态。')
    p('复现入口：scripts/repair_v15/audit_terminal_modifications.py。附表包含完整末端类别计数、10,929条原始条目末端溯源、7,217条有效观察末端表、四组合计数和名称识别探针。证据JSON保存输入哈希、统计、图缓存核对结果。')

    doc=Document();style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    md=[];texts=[]
    for kind, val in blocks:
        if kind=='h': doc.add_heading(val,level=1);md.extend(['# '+val,'']);texts.append(val)
        elif kind=='p': doc.add_paragraph(val);md.extend([val,'']);texts.append(val)
        else:
            t=doc.add_table(rows=0,cols=len(val[0]));t.style='Table Grid'
            for row in val:
                for cell, text in zip(t.add_row().cells,row):cell.text=text;texts.append(text)
            md.extend(['| '+' | '.join(val[0])+' |','| '+' | '.join(['---']*len(val[0]))+' |'])
            md.extend('| '+' | '.join(row)+' |' for row in val[1:]);md.append('')
    basepath=out/'terminal_modification_audit_ZH'
    basepath.with_suffix('.md').write_text('\n'.join(md));doc.save(basepath.with_suffix('.docx'))
    check=Document(basepath.with_suffix('.docx'))
    content='\n'.join([p.text for p in check.paragraphs]+[c.text for t in check.tables for row in t.rows for c in row.cells])
    assert all(text in content for text in texts)
    assert all(sha(path)==value for path,value in inputs.items())
    evidence=dict(status='AUDIT_COMPLETE_TWO_ANNOTATIONS_REQUIRE_REVIEW', timestamp=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(timespec='seconds'),
                  counts=totals, raw_source_ids_matched=len(ids), raw_terminal_mismatch_count=len(mismatches),
                  mismatched_ids=mismatches.dbaasp_id.tolist(), both_mismatches_in_effective_TIGER=True,
                  effective_rows_match_graph_cache=True, cache_path=str(cachepath), training_preprocess_stats=stats,
                  production_or_queue_modified=False, report_text_items_verified=len(texts), input_hashes=inputs)
    dump(out/'terminal_audit_evidence.json',evidence)
    print(json.dumps({'counts':totals,'effective_joint':joint.to_dict('records'),'status':evidence['status'],'output':str(out)},ensure_ascii=False,indent=2))


if __name__ == '__main__':
    main()
