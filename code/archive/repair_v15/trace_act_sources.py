"""Read-only ACT annotation/feature/label provenance; writes a review package only."""
import csv
import importlib.util
import json
from datetime import datetime
from zoneinfo import ZoneInfo
from pathlib import Path
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl import load_workbook
from common import ROOT, WORK, REPORT, ARCHIVE, sha, dump, load
from code.utils.features import calculate_property

OUT=REPORT/'ACT_source_trace_20261012'
HIST=ROOT/'reports/code_methodology_review_20261011/ACT_training_input_audit'
RAW=Path('/data4T/ubuntu/wangyue/postdoc_2025/POAP/Data/dbaasp_jsons')
EXTRACT=RAW.parent/'extract_dbaasp_amp_training_csv.py'
UPSTREAM=Path('/data4T/ubuntu/wangyue/postdoc_2025/TIGER_3D/newdata/dbaasp_amp_training_by_cfu_group_ug_per_mL _20260714.csv')
SOURCE=ROOT/'metadata/train_val_by_cfu_group_ug_per_mL.csv'


def csv_locations(path,encoding='utf-8'):
    out=[]
    with path.open(encoding=encoding,newline='') as f:
        reader=csv.reader(f);next(reader);previous=reader.line_num
        for index,row in enumerate(reader):
            out.append({'source_record_index_0based':index,'source_record_number_with_header':index+2,
                        'source_physical_line_start':previous+1,'source_physical_line_end':reader.line_num})
            previous=reader.line_num
    return pd.DataFrame(out)


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    inputs={}
    def track(path):inputs[str(path)]=sha(path);return path
    for p in [SOURCE,UPSTREAM,EXTRACT,HIST/'ACT_209_training_records_and_features.csv',HIST/'ACT_all_source_records.csv',HIST/'ACT_audit_evidence.json',ROOT/'code/utils/features.py']:
        track(p)
    archived=ARCHIVE/'metadata/train_val_by_cfu_group_ug_per_mL.csv'
    packaged=ROOT/'reports/methods_package_S4_v1/1_training_data_pairing/data/train_val_by_cfu_group_ug_per_mL.csv'
    assert sha(track(archived))==sha(SOURCE)==sha(track(packaged))
    audit=pd.read_csv(HIST/'ACT_209_training_records_and_features.csv')
    source=pd.read_csv(SOURCE);locations=csv_locations(SOURCE)
    assert len(audit)==209 and audit.dbaasp_id.nunique()==audit.sequence.nunique()==204
    selected=source.iloc[audit._source_record_index_0based.astype(int)].reset_index(drop=True)
    for col in ['dbaasp_id','sequence','n_terminus','c_terminus','cfu_group','MIC_Escherichia_coli']:
        assert selected[col].fillna('').astype(str).tolist()==audit[col].fillna('').astype(str).tolist(),col
    upstream=pd.read_csv(UPSTREAM,encoding='gb18030');upstream['upstream_record_number_with_header']=np.arange(len(upstream))+2
    upstream['cfu_key']=upstream.cfu_group.fillna('');audit['cfu_key']=audit.cfu_group.fillna('')
    up=audit.merge(upstream[['dbaasp_id','cfu_key','sequence','n_terminus','c_terminus','MIC_Escherichia_coli','upstream_record_number_with_header']],
                   on=['dbaasp_id','cfu_key'],how='left',validate='one_to_one',suffixes=('','_upstream'))
    for col in ['sequence','n_terminus','c_terminus','MIC_Escherichia_coli']:
        assert up[col].fillna('').astype(str).tolist()==up[col+'_upstream'].fillna('').astype(str).tolist(),col
    spec=importlib.util.spec_from_file_location('act_readonly_extractor',EXTRACT);extract=importlib.util.module_from_spec(spec);spec.loader.exec_module(extract)
    records=[];counter_records=[];peptides=[];activities=[]
    for r in audit.drop_duplicates('dbaasp_id').itertuples():
        path=track(RAW/(r.dbaasp_id.split('_')[-1]+'.json'));j=load(path)
        assert j['dbaaspId']==r.dbaasp_id and j['sequence']==r.sequence
        n=j.get('nTerminus') or {};c=j.get('cTerminus') or {}
        assert n.get('name')=='ACT' and n.get('description')=='Acetyl group; CH3CO'
        assert c.get('name','')==('' if pd.isna(r.c_terminus) else r.c_terminus)
        mw=extract.peptide_mw(j);counter_mw=extract.sequence_mw(j['sequence'],'acetyl',c.get('name',''))
        assert np.isclose(counter_mw-mw,42.0367,rtol=0,atol=1e-9)
        records.append((path,j,mw));counter_records.append((path,j,counter_mw))
        articles=j.get('articles') or []
        peptides.append({'dbaasp_id':r.dbaasp_id,'name':j.get('name',''),'sequence':r.sequence,'length':len(r.sequence),
                         'raw_n_name':n.get('name',''),'raw_n_description':n.get('description',''),
                         'raw_c_name':c.get('name',''),'raw_c_description':c.get('description',''),
                         'raw_c_is_null':j.get('cTerminus') is None,'raw_json_path':str(path),
                         'pubmed_ids':';'.join(str((a.get('pubmed') or {}).get('pubmedId','')) for a in articles),
                         'article_titles':' | '.join(str(a.get('title','')) for a in articles),
                         'raw_has_additional_bonds_or_unusual_residues':any(j.get(k) for k in ['intrachainBonds','interchainBonds','coordinationBonds','unusualAminoAcids']),
                         'mw_ACT_unrecognized':mw,'mw_ACT_recognized_counterfactual':counter_mw,'annotation_check':'MATCH'})
        for ai,a in enumerate(j.get('targetActivities') or []):
            if extract.species_key(extract.nested_name(a.get('targetSpecies')))!='Escherichia coli':continue
            unit=extract.normalize_unit(extract.nested_name(a.get('unit')));value=extract.parse_numeric(a.get('concentration'))
            accepted=extract.is_mic(a) and bool(unit) and value is not None
            activities.append({'dbaasp_id':r.dbaasp_id,'raw_activity_index_0based':ai,'activity_id':a.get('id'),
                               'species_or_strain':extract.nested_name(a.get('targetSpecies')),
                               'endpoint_group':extract.nested_name(a.get('activityMeasureGroup')),'endpoint_value':a.get('activityMeasureValue'),
                               'raw_concentration':a.get('concentration'),'raw_unit':extract.nested_name(a.get('unit')),
                               'parsed_numeric_existing_rule':value,'normalized_unit':unit,'cfu_group':extract.nested_name(a.get('cfuGroup')),
                               'cfu':a.get('cfu'),'medium':extract.nested_name(a.get('medium')),'ph':a.get('ph'),
                               'reference_index':a.get('reference'),'raw_activity_field':a.get('activity'),
                               'included_by_existing_extractor':accepted,
                               'mass_MIC_existing_rule':extract.converted_value(value,unit,'ug_per_mL',mw) if accepted else None,
                               'mass_MIC_ACT_counterfactual':extract.converted_value(value,unit,'ug_per_mL',counter_mw) if accepted else None,
                               'raw_json_path':str(path)})
    old=pd.DataFrame(extract.build_rows(records,[],'ug_per_mL'));counter=pd.DataFrame(extract.build_rows(counter_records,[],'ug_per_mL'))
    for df in [old,counter]:df['cfu_key']=df.cfu_group.fillna('')
    audit=audit.merge(old[['dbaasp_id','cfu_key','MIC_Escherichia_coli']].rename(columns={'MIC_Escherichia_coli':'mic_rebuilt_existing_extractor'}),on=['dbaasp_id','cfu_key'],validate='one_to_one')
    audit=audit.merge(counter[['dbaasp_id','cfu_key','MIC_Escherichia_coli']].rename(columns={'MIC_Escherichia_coli':'mic_ACT_mass_counterfactual'}),on=['dbaasp_id','cfu_key'],validate='one_to_one')
    for col in ['mic_rebuilt_existing_extractor','mic_ACT_mass_counterfactual']:
        audit[col]=pd.to_numeric(audit[col],errors='raise').astype(float)
    assert (audit.MIC_Escherichia_coli==audit.mic_rebuilt_existing_extractor).all()
    audit['mic_delta_counterfactual']=audit.mic_ACT_mass_counterfactual-audit.MIC_Escherichia_coli
    audit['mic_relative_change_percent']=100*audit.mic_delta_counterfactual/audit.MIC_Escherichia_coli
    audit['log2_mic_delta_counterfactual']=np.log2(audit.mic_ACT_mass_counterfactual/audit.MIC_Escherichia_coli)
    audit['MIC_changed_by_ACT_mass']=audit.mic_delta_counterfactual.abs()>1e-6
    new=np.array([calculate_property(r.sequence,r.n_terminus,r.c_terminus) for r in audit.itertuples()])
    assert np.allclose(new[:,7],audit.charge_ACT_as_acetyl,rtol=0,atol=1e-6)
    assert np.allclose(new[:,9],audit.charge_density_ACT_as_acetyl,rtol=0,atol=1e-6)
    audit['charge_current']=new[:,7];audit['charge_density_current']=new[:,9]
    audit=audit.merge(locations,left_on='_source_record_index_0based',right_on='source_record_index_0based',validate='one_to_one')
    audit=audit.merge(up[['dbaasp_id','cfu_key','upstream_record_number_with_header']],on=['dbaasp_id','cfu_key'],validate='one_to_one')
    pep=pd.DataFrame(peptides);act=pd.DataFrame(activities)
    accepted=act[act.included_by_existing_extractor]
    audit['raw_MIC_units_same_cfu']=[';'.join(sorted(accepted.loc[(accepted.dbaasp_id==r.dbaasp_id)&(accepted.cfu_group.fillna('')==r.cfu_key),'raw_unit'].dropna().unique())) for r in audit.itertuples()]
    audit=audit.merge(pep[['dbaasp_id','raw_n_description','raw_c_is_null','raw_json_path','mw_ACT_unrecognized','mw_ACT_recognized_counterfactual']],on='dbaasp_id',validate='many_to_one')
    audit['source_csv_path']=str(SOURCE);audit['upstream_csv_path']=str(UPSTREAM)
    audit['review_decision']='';audit['review_notes']=''
    first=['source_record_number_with_header','source_physical_line_start','dbaasp_id','name','sequence','length','n_terminus','c_terminus','normalized_cfu','MIC_Escherichia_coli','charge_old','charge_current','charge_density_old','charge_density_current','MIC_changed_by_ACT_mass','mic_ACT_mass_counterfactual','mic_delta_counterfactual','mic_relative_change_percent','raw_MIC_units_same_cfu','pubmed_ids','raw_json_path','upstream_record_number_with_header','review_decision','review_notes']
    audit=audit[first+[c for c in audit if c not in first]]
    all_source=pd.read_csv(HIST/'ACT_all_source_records.csv').merge(locations,left_on='_source_record_index_0based',right_on='source_record_index_0based',validate='one_to_one')
    all_source['source_csv_path']=str(SOURCE)
    missing=audit[audit.c_terminus.isna()];multiple=audit[audit.sequence.duplicated(keep=False)];changed=audit[audit.MIC_changed_by_ACT_mass]
    raw_selected=selected.copy();raw_selected.insert(0,'source_record_number_with_header',audit.source_record_number_with_header.to_numpy())
    tables={'ACT_209_training_trace':audit,'ACT_204_raw_peptides':pep,'ACT_138_MIC_mass_changes':changed,
            'ACT_raw_Ecoli_activities':act,'ACT_399_source_disposition':all_source,
            'ACT_5_missing_C_terminus':missing,'ACT_5_multicfu_peptides_10rows':multiple,'ACT_209_original_source_columns':raw_selected}
    for name,df in tables.items():df.to_csv(OUT/f'{name}.csv',index=False,encoding='utf-8-sig')
    affected=changed.mic_relative_change_percent
    evidence={'status':'AUDIT_COMPLETE_ADDITIONAL_LABEL_BUG_CONFIRMED','created_at':datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(timespec='seconds'),
              'effective_training_observations':7217,'ACT_raw_rows':399,'ACT_training_observations':209,'ACT_unique_sequences_and_ids':204,
              'raw_json_annotation_matches':204,'upstream_csv_exact_matches':209,'source_labels_reproduced_exactly':209,
              'current_feature_repair_verified':209,'missing_C_terminus_observations':5,
              'ACT_label_counterfactual_changed':len(changed),'ACT_label_counterfactual_unchanged':len(audit)-len(changed),
              'affected_label_relative_percent_min':float(affected.min()),'affected_label_relative_percent_max':float(affected.max()),
              'affected_label_relative_percent_median':float(affected.median()),'counterfactual_acetyl_mass_increment':42.0367,
              'scope':'209 currently selected ACT observations only; fixed input records and existing extraction/aggregation rules, no new deduplication',
              'production_inputs_modified':False,'training_tasks_modified':False,'sources_sha256':inputs}
    dump(OUT/'ACT_source_trace_evidence.json',evidence)
    glossary=[['字段','含义'],['source_record_number_with_header','原训练CSV的记录序号，表头计第1条；不是绝对保证的文本行号。'],
              ['source_physical_line_start/end','原训练CSV物理文本行起止；按CSV解析器计算，支持带换行的字段。'],
              ['upstream_record_number_with_header','20260714上游CFU表的记录序号，表头计第1条；源文件为gb18030编码。'],
              ['training_row_index_0based','7,217条有效训练观察中的0起始行索引。'],['charge_old / charge_current','旧/现行特征函数的原始电荷描述符；不是实验测量电荷。'],
              ['mic_ACT_mass_counterfactual','仅修复ACT分子量别名后的反事实MIC，未写回训练表，尚非完整重新清洗后的正式标签。'],
              ['raw_concentration / raw_unit','DBAASP JSON中的原始活性浓度字符串和单位；不等于CSV内的物种/CFU组聚合标签。'],
              ['included_by_existing_extractor','按既有端点、单位和数值解析规则是否参与提取；沿用旧规则便于隔离ACT质量影响。'],
              ['raw_activity_index_0based / activity_id','JSON targetActivities数组的0起始位置及数据库活动ID。'],
              ['raw_c_is_null','原JSON的cTerminus是否为null；null不作为已证实自由C端。'],
              ['review_decision / review_notes','预留给用户填写的核查结论和备注。']]
    notes=[['事项','说明'],['目的','定位209条训练观察及204条原始肽，核对ACT注释、特征输入、MIC换算；不改训练数据。'],
           ['核心结论','204份原始JSON均明确ACT=Acetyl group; CH3CO，标注传递一致。旧训练特征函数漏识别ACT。'],
           ['新增发现','上游MIC分子量函数也漏识别ACT。209条现标签全部精确复现；仅修正别名后138条会变化，71条不变。'],
           ['计数区别','209条sequence×CFU观察，204条不同序列/DBAASP ID，5条肽各有两条CFU观察。'],
           ['需原文核查','5条C端null记录；原文实验端点与原始浓度可通过活动表及PubMed编号定位。'],
           ['修复边界','当前重训修复了理化特征，尚未修正这里新发现的MIC标签换算问题；现有重训结果应视作阶段性结果。'],
           ['记录序号','Excel原CSV导入后对应记录序号；物理行号另列，不混用。'],
           ['反事实范围','保持现有入选记录及提取/均值规则固定；没有全训练池重筛选、跨方法传播评估或再次训练。']]
    xlsx=OUT/'ACT_training_source_review.xlsx'
    sheets={'说明':pd.DataFrame(notes[1:],columns=notes[0]),'209训练观察':audit,'204原始肽':pep,'138_MIC换算差异':changed,
            '原始Ecoli活性':act,'399源记录去向':all_source,'5_C端缺失':missing,'5_多CFU肽':multiple,'209原CSV全字段':raw_selected,
            '字段说明':pd.DataFrame(glossary[1:],columns=glossary[0])}
    with pd.ExcelWriter(xlsx,engine='openpyxl') as writer:
        for name,df in sheets.items():df.to_excel(writer,sheet_name=name,index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes='E2' if ws.max_column>8 else 'A2';ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='245A78');cell.alignment=Alignment(wrap_text=True,vertical='top')
            ws.row_dimensions[1].height=42
            for col in ws.columns:
                header=col[0].value;letter=col[0].column_letter
                ws.column_dimensions[letter].width=65 if header in ['说明','含义','raw_json_path','article_titles'] else 48 if header=='sequence' else 28 if header=='name' else 23
                for cell in col[1:]:
                    cell.alignment=Alignment(vertical='top',wrap_text=header in ['说明','含义'])
                    if header=='raw_json_path' and cell.value:cell.hyperlink=Path(cell.value).as_uri();cell.font=Font(color='0563C1',underline='single')
            ws.sheet_properties.pageSetUpPr.fitToPage=True
    wb=load_workbook(xlsx,read_only=True,data_only=True)
    assert wb['209训练观察'].max_row==210 and wb['204原始肽'].max_row==205 and wb['138_MIC换算差异'].max_row==139
    wb.close()
    blocks=[]
    def h(t):blocks.append(('h',t))
    def p(t):blocks.append(('p',t))
    def table(t):blocks.append(('t',t))
    h('ACT训练肽原始记录溯源与新增MIC换算问题')
    p('核查时间：'+evidence['created_at']+'。本次为只读溯源和反事实复算；没有修改源数据、特征函数、训练任务或已生成模型。提供Excel、CSV及同内容Markdown/Word，便于逐条核查。')
    p('结论一：204条不同肽的原始DBAASP JSON均写明nTerminus.name=ACT、description=Acetyl group; CH3CO；序列、N/C端字段与实际训练源表一致。因此，本批ACT注释没有发现从原JSON到训练表的转录错误，已确认的问题是旧理化特征函数漏识别ACT。这里验证的是本地数据库导出与训练表的一致性，不等于已逐篇核对原论文或实验证明化学结构。')
    p('结论二（本次新增）：上游数据提取程序的分子量函数同样只识别ace/acetyl，漏掉ACT。它在µM转µg/mL时遗漏乙酰基质量，确实涉及标签整理。用原始204份JSON和该提取程序，可逐条精确复现209条现有E. coli训练标签；仅将ACT映射为acetyl后，其中138条标签增加，71条不变。当前重训已修复理化特征，但使用的训练CSV尚未修复此标签换算问题，现有重训结果仍是阶段性结果。')
    h('一、受影响肽及实际输入范围')
    table([['范围','数量/说明'],['训练源表','11,516条记录，其中399条N端ACT'],['MIC缺失或非正被排除','154条ACT记录'],['长度不在6–50 aa被排除','6条ACT记录'],['去重未被选中','30条ACT记录'],['实际进入训练','209条sequence×CFU观察；占7,217条观察的2.8959%'],['不同肽','204条不同序列，204个DBAASP ID'],['重复观察解释','5条肽各出现于两个CFU组，共10条观察；不是209条不同肽'],['C端','204条观察AMD，5条缺失；按不同肽计199条AMD、5条缺失'],['长度','6–50 aa，中位数15 aa'],['CFU','1E5–1E6：119；1E4–1E5：22；1E6–1E7：8；unknown：60']])
    p('修复前后训练源CSV与包内CSV、修复前归档SHA-256相同。已核对所选209条记录与20260714上游CFU表的序列、末端及MIC逐条一致；原始JSON中未发现这204条肽带额外的intrachain/interchain/coordination bonds或unusualAminoAcids条目，但未用此字段检查替代独立化学结构验证。')
    h('二、在哪里查原始文件')
    table([['层次','文件或目录'],['实际训练源表',str(SOURCE)],['上游CFU表（gb18030编码）',str(UPSTREAM)],['原始DBAASP JSON',str(RAW)+'/<DBAASP ID数字部分>.json'],['上游提取代码',str(EXTRACT)],['本轮对照工作簿',str(xlsx)]])
    p('工作簿逐条提供训练CSV记录序号、物理行起止、上游CSV记录序号、训练行索引、DBAASP ID、名称、完整序列、两端标注、CFU、MIC、原始活性浓度与单位、PubMed ID、JSON路径及特征旧/新值。JSON内重点查看nTerminus、cTerminus、targetActivities和articles。记录序号把表头计为第1条；带换行字段时不能将记录序号误当文本物理行号。')
    h('三、训练特征到底改了什么')
    p('旧函数将ACT转成小写act后未命中N端封闭别名，按未封闭N端增加1；修复后把act映射到acetyl。原始10维理化向量只直接改变第8维charge和第10维charge_density：旧值分别高出1和约1/L（L为序列长度，落盘值经过三位小数舍入）。其余描述符维持既有定义。这里的电荷差是代码内的描述符差值，不是实验测量误差。')
    examples=audit.head(4)
    table([['CSV记录号','DBAASP ID/名称','序列','charge旧→现行','density旧→现行']]+[[str(int(r.source_record_number_with_header)),r.dbaasp_id+' / '+str(r.name),r.sequence,f'{r.charge_old:.3f} → {r.charge_current:.3f}',f'{r.charge_density_old:.3f} → {r.charge_density_current:.3f}'] for r in examples.itertuples()])
    p('所有209条当前特征均与修复前专项审计保存的ACT→acetyl反事实值一致。各折训练集重新拟合charge和charge_density的均值/标准差，因此其他未带ACT的肽也会有这两列标准化数值变化；序列、结构、末端原始标注及MIC不因这一特征代码修复自动改变。尤其不能把“特征已修复”等同于“上游MIC转换已修复”。')
    h('四、138条MIC标签换算的新增证据')
    p('上游sequence_mw先按残基质量加水计算分子量，识别acetyl时加42.0367 Da，识别AMD时减0.9840 Da；ACT没有命中N端分支。现有converted_value使用MIC(µg/mL)=MIC(µM)×MW/1000。下表保持原始活动筛选、数值解析、同物种/CFU组算术平均及入选记录不变，仅修复ACT别名；它不是重新清洗后的最终训练标签。')
    table([['DBAASP ID','原CSV记录号','现标签 µg/mL','ACT质量反事实 µg/mL','变化%']]+[[r.dbaasp_id,str(int(r.source_record_number_with_header)),f'{r.MIC_Escherichia_coli:.6f}',f'{r.mic_ACT_mass_counterfactual:.6f}',f'{r.mic_relative_change_percent:+.3f}'] for r in changed.head(6).itertuples()])
    p(f'138条变化标签的相对增幅为{affected.min():.3f}%–{affected.max():.3f}%，中位数{affected.median():.3f}%。直接以µg/mL报告的记录不需要分子量换算；同组若混有µM和µg/mL，只有前者的转换分量改变。这个计数仅覆盖当前入选209条ACT观察，尚未重新评估其他末端别名、全训练池去重选择或外部基准训练表的传播影响。')
    p('具体示例：DBAASPS_9784（BM0，RFWWFRRR）原JSON活动71579为E. coli DH5alpha、MIC=8 µM；当前提取函数得到MW=1308.56088，标签10.468487 µg/mL。按该脚本既有乙酰基常数补加42.0367，得到MW=1350.59758，标签10.804781 µg/mL。旧值能够精确复现，因此不是单纯比较两个年代CSV时出现的未知版本差异。该反事实沿用脚本质量常数，未把数据库activity字段的末位差异强行归为同一个值。')
    h('五、建议优先人工核对的记录')
    p('优先检查下列5条C端缺失记录：原JSON的cTerminus也为null，并非本次CSV整理丢失。现有特征计算对空值沿用默认自由C端处理，但缺失字段不能证明实际为自由C端；建议查对应论文的合成方法或序列表。')
    table([['DBAASP ID','名称','完整序列','CSV记录号','PubMed']]+[[r.dbaasp_id,'' if pd.isna(r.name) else str(r.name),r.sequence,str(int(r.source_record_number_with_header)),str(r.pubmed_ids)] for r in missing.itertuples()])
    p('另有5条肽因不同CFU组形成两条观察，逐条MIC和源行号保存在5_多CFU肽工作表：Pac-525（DBAASPS_710）、V681[V13K,T15K]（DBAASPS_4975）、V681[V13K]/Peptide P（DBAASPS_265）、Histatin 5(4–15)/P-113（DBAASPS_16041）、1-RW（DBAASPS_10875）。这解释209观察与204序列的差异，本身不等同于重复整理错误。')
    p('核查时可按ID打开JSON，先确认乙酰化/酰胺化，再按CFU及activity_id追到原始E. coli活性，核对浓度单位、MIC端点、菌株与文献。当前提取程序会处理区间和不等号、并按物种和CFU组聚合；工作簿保留原始浓度字符串和菌株，方便人工检查。ACT反事实沿用这些旧规则，不表示本次已为所有聚合/删失处理作方法学背书。')
    h('六、交付与结论边界')
    table([['交付','内容'],['ACT_training_source_review.xlsx','10个工作表，可筛选；含预留核查结论/备注列'],['ACT_209_training_trace.csv','全部209条受影响训练观察及源定位、特征、标签反事实'],['ACT_204_raw_peptides.csv','204份JSON的末端定义、文章、质量及路径'],['ACT_138_MIC_mass_changes.csv','本次新确认的标签换算差异'],['ACT_raw_Ecoli_activities.csv','204条肽的原始E. coli活动字段与既有提取规则对照'],['ACT_399_source_disposition.csv','所有399条ACT源记录的过滤/去重去向'],['ACT_source_trace_evidence.json','核验结果、范围限制及输入SHA-256'],['scripts/repair_v15/trace_act_sources.py','本次只读复算及报告生成代码']])
    p('目前可确认：ACT标注从本地原始JSON到上游表、实际训练表传递一致；特征函数漏识别ACT已修复；上游MIC单位换算另有同类漏识别，尚未写回生产数据。请将当前模型比较视为基于现有标签的阶段性结果。全量标签修复及其训练依赖范围需另行核定，本报告不把138条反事实数值直接写成最终训练集。')
    p('当前GPU队列继续执行原已批准任务；本次溯源没有取消或重启作业，也没有追加重训。历史2026-10-11专项报告描述的是修复前时点，本报告补充原始JSON溯源与新发现，不覆盖历史证据。')
    doc=Document();doc.sections[0].top_margin=Inches(.6);doc.sections[0].bottom_margin=Inches(.6)
    doc.sections[0].left_margin=Inches(.65);doc.sections[0].right_margin=Inches(.65)
    style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(9);style.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    md=[];texts=[]
    for kind,val in blocks:
        if kind=='h':doc.add_heading(val,level=1);md.extend(['# '+val,'']);texts.append(val)
        elif kind=='p':doc.add_paragraph(val);md.extend([val,'']);texts.append(val)
        else:
            t=doc.add_table(rows=0,cols=len(val[0]));t.style='Table Grid'
            for row in val:
                for cell,text in zip(t.add_row().cells,row):cell.text=text;texts.append(text)
            md.extend(['| '+' | '.join(val[0])+' |','| '+' | '.join(['---']*len(val[0]))+' |'])
            md.extend('| '+' | '.join(row)+' |' for row in val[1:]);md.append('')
    base=OUT/'ACT_source_trace_report_ZH';base.with_suffix('.md').write_text('\n'.join(md));doc.save(base.with_suffix('.docx'))
    checked=Document(base.with_suffix('.docx'));actual='\n'.join([p.text for p in checked.paragraphs]+[c.text for t in checked.tables for row in t.rows for c in row.cells])
    assert all(t in actual for t in texts)
    assert all(sha(Path(path))==value for path,value in inputs.items())
    dump(OUT/'delivery_validation.json',{'status':'PASS','source_hashes_unchanged':len(inputs),'report_strings_matched':len(texts),
                                       'xlsx_training_rows':209,'xlsx_raw_peptides':204,'xlsx_MIC_difference_rows':138,
                                       'output_hashes':{p.name:sha(p) for p in OUT.iterdir() if p.is_file() and p.name!='delivery_validation.json'}})
    print(json.dumps({k:v for k,v in evidence.items() if k!='sources_sha256'},ensure_ascii=False,indent=2));print(OUT)


if __name__=='__main__':main()
