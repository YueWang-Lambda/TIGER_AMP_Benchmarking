"""Audit completed S4/P10 seeds against immutable pre-repair prediction hashes."""
from datetime import datetime
from zoneinfo import ZoneInfo
import numpy as np
import pandas as pd
from scipy.stats import pearsonr, kendalltau, wilcoxon
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from common import ROOT, WORK, REPORT, ARCHIVE, load, dump, sha
from external_metrics import evaluate, METRICS


def main():
    oldroot=ROOT/'reports/S4_study_nature_v14'
    manifest=ARCHIVE/'reports/S4_study_nature_v14/rse_prediction_replay_manifest.json'
    old_refs={int(r['seed']):r for r in load(manifest)['predictions'] if r['method']=='TIGER-pair'}
    assert set(old_refs)==set(range(120,130))
    archived_metrics=load(ARCHIVE/'reports/methods_package_S4_v1/3_S4_ablation/results/s4_panels_fourmetrics.json')['S4']['per_seed']
    rows=[]; families=[]; sources=[]
    for seed in range(120,130):
        newdir=WORK/f'tiger_external/S4_seed{seed}'
        assert load(newdir/'EVALUATED.json')['status']=='PASS'
        assert load(ROOT/f'runs_repair_v15/S4/S4_seed{seed}/DONE.json')['status']=='PASS'
        oldpath=oldroot/old_refs[seed]['cache_file']
        assert sha(oldpath)==old_refs[seed]['sha256'],('Old prediction modified',seed)
        for version,path in [('old',oldpath),('new',newdir/'pair_predictions.csv')]:
            met,fam,ordered=evaluate(pd.read_csv(path))
            assert met['n_pairs']==2211 and met['n_families']==11
            assert all(not values for values in met['undefined_families'].values())
            # Independent direct formulas check the common metric implementation.
            direct=[]
            for family,g in ordered.groupby('family'):
                y=g.y_true.to_numpy(float); p=g.y_pred.to_numpy(float)
                direct.append({'PCC':pearsonr(y,p).statistic,'KCC':kendalltau(y,p).statistic,
                               'log10MAE':np.mean(np.abs(y-p))*np.log10(2),
                               'RSE':np.sum((y-p)**2)/np.sum((y-y.mean())**2)})
            assert np.allclose(pd.DataFrame(direct)[METRICS].mean().to_numpy(),
                               [met['macro'][k] for k in METRICS],rtol=0,atol=1e-12)
            if version=='old':
                assert all(abs(met['macro'][k]-archived_metrics[str(seed)][k])<5.1e-5 for k in METRICS)
            else:
                stored=load(newdir/'metrics.json')['macro']
                assert all(abs(met['macro'][k]-stored[k])<1e-12 for k in METRICS)
            rows.append({'version':version,'seed':seed,**met['macro']})
            families.append(fam.assign(version=version,seed=seed))
            sources.append({'version':version,'seed':seed,'path':str(path.relative_to(ROOT)),
                            'sha256':sha(path),'n_pairs':2211,'n_families':11})
    data=pd.DataFrame(rows); family_data=pd.concat(families,ignore_index=True)
    old=data[data.version=='old'].set_index('seed'); new=data[data.version=='new'].set_index('seed')
    summary=[]
    for k in METRICS:
        a=old[k].to_numpy(); b=new[k].to_numpy(); d=b-a
        summary.append({'metric':k,'old_mean':float(a.mean()),'old_sd_ddof0':float(a.std(ddof=0)),
                        'new_mean':float(b.mean()),'new_sd_ddof0':float(b.std(ddof=0)),
                        'delta_new_minus_old':float(d.mean()),'wilcoxon_two_sided_exact_p':float(wilcoxon(d,method='exact').pvalue),
                        'n_improved':int(np.sum(d>0 if k in ['PCC','KCC'] else d<0)),
                        'old_sd_ddof1':float(a.std(ddof=1)),'new_sd_ddof1':float(b.std(ddof=1))})
    order=np.argsort([r['wilcoxon_two_sided_exact_p'] for r in summary]); maximum=0.
    for rank,i in enumerate(order):
        maximum=max(maximum,min(1.,(len(summary)-rank)*summary[i]['wilcoxon_two_sided_exact_p']))
        summary[i]['holm_p_four_metrics']=maximum
    family_mean=family_data.groupby(['family','version'])[METRICS].mean().unstack('version')
    family_rows=[]
    for family in family_mean.index:
        row={'family':family}
        for k in METRICS:
            row[f'{k}_old']=float(family_mean.loc[family,(k,'old')]);row[f'{k}_new']=float(family_mean.loc[family,(k,'new')])
            row[f'{k}_delta']=row[f'{k}_new']-row[f'{k}_old']
        family_rows.append(row)
    REPORT.mkdir(parents=True,exist_ok=True); out=WORK/'tiger_baseline_comparison';out.mkdir(exist_ok=True)
    data.to_csv(out/'per_seed.csv',index=False);family_data.to_csv(out/'family_per_seed.csv',index=False)
    pd.DataFrame(family_rows).to_csv(out/'family_comparison.csv',index=False)
    pd.DataFrame(summary).to_csv(out/'summary.csv',index=False)
    timestamp=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(timespec='seconds')
    dump(out/'comparison.json',{'status':'PASS','created_at':timestamp,'comparison':'S4/P10 pre-repair pipeline versus repaired pipeline',
                              'n_seeds':10,'seed_ids':list(range(120,130)),'sd_display_ddof':0,
                              'summary':summary,'sources':sources,'old_manifest_sha256':sha(manifest),
                              'independent_metric_formula_check':True})
    blocks=[]
    def p(text):blocks.append(('p',text))
    def table(rows):blocks.append(('table',rows))
    blocks.append(('h','TIGER-pair完整主模型：重训前后外测对比'))
    p('核查时间：'+timestamp+'。本轮完整主模型10个种子（120–129）的五折训练和外测均已完成。与修复前相比，宏PCC由0.6313降至0.6200，RSE由0.6720升至0.6798，点估计均小幅变差。')
    p('比较对象为原S4完整主模型，即负样本分解中的P10：分类预训练使用负样本，回归阶段不注入负样本。它不是P00纯无负样本组，也不是P11两阶段均使用负样本组；此处不以P11替换原主模型或按外测结果选择配置。')
    p('统一口径：同一11家族、208肽、2,211条有向边，逐家族计算四指标后等权宏平均，再对10种子汇总；每种子采用五折集成及其匹配的校准器。表中±为种子总体标准差（ddof=0），沿用旧报告口径，非标准误；另在机器可读结果中保留ddof=1标准差。旧预测10份文件均与修复前归档manifest的SHA-256一致，旧指标复算与归档四舍五入值一致；新预测复算与本轮metrics.json一致，直接公式二次校验通过。')
    p('全部统计由未舍入的逐边预测计算，最后统一格式化；因此旧PCC标准差显示为0.0253，与历史汇总表的0.0252存在末位舍入差异，不代表旧预测或模型发生变化。')
    table([['指标','重训前：均值±SD','本轮：均值±SD','变化（新−旧）','配对Wilcoxon p','Holm校正p']]+[
        [r['metric']+(' ↑' if r['metric'] in ['PCC','KCC'] else ' ↓'),f"{r['old_mean']:.4f} ± {r['old_sd_ddof0']:.4f}",
         f"{r['new_mean']:.4f} ± {r['new_sd_ddof0']:.4f}",f"{r['delta_new_minus_old']:+.4f}",
         f"{r['wilcoxon_two_sided_exact_p']:.4f}",f"{r['holm_p_four_metrics']:.4f}"] for r in summary])
    p('PCC为3个种子升高、7个降低；RSE为3个种子改善、7个变差。四项指标的双侧精确Wilcoxon检验均未达到0.05；按四指标作Holm校正后也未达到。检验单位为相同编号的10个训练种子，仅描述固定外测集上的训练波动；不把2,211条共享端点的边当成独立重复，也不能据未检出差异宣称新旧等效。')
    table([['种子','旧PCC','新PCC','PCC变化','旧RSE','新RSE']]+[
        [str(seed),f"{old.loc[seed,'PCC']:.4f}",f"{new.loc[seed,'PCC']:.4f}",
         f"{new.loc[seed,'PCC']-old.loc[seed,'PCC']:+.4f}",f"{old.loc[seed,'RSE']:.4f}",f"{new.loc[seed,'RSE']:.4f}"] for seed in range(120,130)])
    p('家族变化并不一致：PCC下降最大的为Repeat（−0.0646）、Arctoterin-1（−0.0400）和Mammuthusin-2（−0.0352）；Hesperelin-3（+0.0249）和Mammuthusin-3（+0.0180）有所提高。以下为10种子的家族指标均值，属于描述性分解，未进行家族优劣的确认性检验。')
    table([['家族','旧PCC','新PCC','PCC变化','旧RSE','新RSE']]+[
        [r['family'],f"{r['PCC_old']:.4f}",f"{r['PCC_new']:.4f}",f"{r['PCC_delta']:+.4f}",f"{r['RSE_old']:.4f}",f"{r['RSE_new']:.4f}"] for r in family_rows])
    p('解释边界：这里比较修复前与修复后的整条流程，涉及ACT训练理化特征/标准化及外测CFU等输入规范化，未单独隔离各项修复的贡献，不能把变化全部归因于ACT。保持外测标签、边集合和汇总定义一致，并不意味着推理输入特征完全相同。该结果已覆盖完整主模型10种子，但其余消融和负样本分解尚未全部完成，最终跨方法排名、正式主表和图文仍以v15统一验收为准。')
    label_audit=REPORT/'ACT_source_trace_20261012/ACT_source_trace_evidence.json'
    if label_audit.exists():
        evidence=load(label_audit)
        p(f"2026-10-12原始记录溯源补记：上游MIC分子量换算程序另有ACT别名遗漏，当前209条ACT训练观察中{evidence['ACT_label_counterfactual_changed']}条MIC在仅修复该遗漏的反事实下发生变化。现有重训使用的CSV尚未修正此标签问题，因此本报告保留为基于现有标签的阶段性前后比较，不作为标签修复后的最终v15结果。详见ACT_source_trace_20261012/ACT_source_trace_report_ZH.md/.docx及逐记录工作簿。")
    p('复算代码：scripts/repair_v15/compare_tiger_baseline.py。机器可读输出：artifacts/repair_v15/tiger_baseline_comparison/，含comparison.json、summary.csv、per_seed.csv、family_per_seed.csv及family_comparison.csv。源预测路径与SHA-256逐份登记于comparison.json；本次未修改训练、推理代码或已有科学结果。')
    doc=Document();style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');md=[];texts=[]
    for kind,val in blocks:
        if kind=='h':doc.add_heading(val,level=1);md.extend(['# '+val,'']);texts.append(val)
        elif kind=='p':doc.add_paragraph(val);md.extend([val,'']);texts.append(val)
        else:
            t=doc.add_table(rows=0,cols=len(val[0]));t.style='Table Grid'
            for row in val:
                for cell,text in zip(t.add_row().cells,row):cell.text=text;texts.append(text)
            md.extend(['| '+' | '.join(val[0])+' |','| '+' | '.join(['---']*len(val[0]))+' |'])
            md.extend('| '+' | '.join(row)+' |' for row in val[1:]);md.append('')
    base=REPORT/'TIGER_pair_retraining_comparison_ZH';base.with_suffix('.md').write_text('\n'.join(md));doc.save(base.with_suffix('.docx'))
    verify=Document(base.with_suffix('.docx'));content='\n'.join([p.text for p in verify.paragraphs]+[c.text for t in verify.tables for row in t.rows for c in row.cells])
    assert all(text in content for text in texts)
    print(pd.DataFrame(summary).to_string(index=False));print(base)


if __name__=='__main__':main()
