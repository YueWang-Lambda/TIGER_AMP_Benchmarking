"""Read-only propagation of the ACT molecular-mass correction into task scope."""
import importlib.util
from datetime import datetime
from zoneinfo import ZoneInfo
import numpy as np
import pandas as pd
from sklearn.model_selection import GroupKFold
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from common import ROOT, WORK, REPORT, load, dump, sha
from external_one import METHODS
from code.utils.features import normalize_cfu_group


def main():
    out=REPORT/'ACT_retraining_scope_20261012';out.mkdir(parents=True,exist_ok=True)
    source_path=ROOT/'metadata/train_val_by_cfu_group_ug_per_mL.csv'
    restricted_path=ROOT/'reports/methods_package_S4_v1/4_S4_benchmark/data/train_restricted4929_tiger.csv'
    rawdir=ROOT.parents[2]/'POAP/Data/dbaasp_jsons'
    script=rawdir.parent/'extract_dbaasp_amp_training_csv.py'
    # ROOT.parents[2] is postdoc_2025 in this repository layout.
    assert rawdir.exists()
    spec=importlib.util.spec_from_file_location('act_scope_extractor',script);m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
    hashes={str(p):sha(p) for p in [source_path,restricted_path,script]}
    old=pd.read_csv(source_path);new=old.copy();act=old[old.n_terminus.fillna('').str.strip().str.upper().eq('ACT')].copy()
    act['source_index']=act.index;act.cfu_group=act.cfu_group.fillna('')
    old_records=[];new_records=[]
    for did in act.dbaasp_id.unique():
        p=rawdir/(did.split('_')[-1]+'.json');j=load(p);hashes[str(p)]=sha(p)
        old_records.append((p,j,m.peptide_mw(j)))
        new_records.append((p,j,m.sequence_mw(j['sequence'],'acetyl',(j.get('cTerminus') or {}).get('name',''))))
    a=pd.DataFrame(m.build_rows(old_records,[],'ug_per_mL'));b=pd.DataFrame(m.build_rows(new_records,[],'ug_per_mL'))
    keys=['dbaasp_id','cfu_group'];target='MIC_Escherichia_coli'
    for df in [a,b]:df.cfu_group=df.cfu_group.fillna('');df[target]=pd.to_numeric(df[target],errors='coerce')
    change=act.merge(a[keys+[target]],on=keys,suffixes=('','_rebuild'),validate='one_to_one')
    change=change.merge(b[keys+[target]],on=keys,suffixes=('','_new'),validate='one_to_one')
    assert len(change)==399 and np.isclose(change[target],change[target+'_rebuild'],rtol=0,atol=1e-6,equal_nan=True).all()
    new.loc[change.source_index,target]=change[target+'_new'].to_numpy()
    cfg=load(ROOT/'runs_repair_v15/S4/S4_seed120/config.json')
    pdbdir=ROOT/cfg['train_pdb_dir'];pdb_names={p.stem for p in pdbdir.glob('*.pdb')}
    def effective(frame):
        frame=frame.copy();frame['source_index']=np.arange(len(frame));frame['sequence_norm']=frame.sequence.str.strip().str.upper()
        frame['cfu_norm']=frame.cfu_group.map(normalize_cfu_group);frame['isupper']=frame.sequence.str.strip().eq(frame.sequence.str.strip().str.upper())
        frame=frame[(frame[target]>0)&frame.sequence_norm.str.fullmatch('[ACDEFGHIKLMNPQRSTVWY]+')&frame.sequence_norm.str.len().between(cfg['min_len'],cfg['max_len'])&frame.sequence_norm.isin(pdb_names)]
        rows=[]
        for _,g in frame.groupby(['sequence_norm','cfu_norm'],sort=False):
            c=g[g.isupper] if g.isupper.any() else g;rows.append(c.iloc[int(np.argmin(c[target].to_numpy()))])
        return pd.DataFrame(rows).reset_index(drop=True)
    x=effective(old);y=effective(new)
    assert len(x)==len(y)==7217 and x[['sequence_norm','cfu_norm']].equals(y[['sequence_norm','cfu_norm']])
    diff=np.abs(y[target]-x[target])>1e-6;representative=x.source_index!=y.source_index
    eff=pd.DataFrame({'sequence':x.sequence_norm,'cfu_group':x.cfu_norm,'old_source_index':x.source_index,'new_source_index':y.source_index,
                      'old_id':x.dbaasp_id,'new_id':y.dbaasp_id,'old_n_terminus':x.n_terminus,'new_n_terminus':y.n_terminus,
                      'old_MIC':x[target],'new_MIC':y[target],'label_changed':diff,'representative_changed':representative})
    eff[diff|representative].to_csv(out/'TIGER_effective_changes.csv',index=False,encoding='utf-8-sig')
    eff[representative].to_csv(out/'TIGER_two_representative_changes.csv',index=False,encoding='utf-8-sig')
    folds=[]
    for fold,(tr,va) in enumerate(GroupKFold(5).split(x,groups=x.sequence_norm),1):
        folds.append({'fold':fold,'changed_train_labels':int(diff.iloc[tr].sum()),'changed_validation_labels':int(diff.iloc[va].sum())})
    assert all(r['changed_train_labels']>0 and r['changed_validation_labels']>0 for r in folds)
    pd.DataFrame(folds).to_csv(out/'TIGER_fold_label_changes.csv',index=False)
    restricted=pd.read_csv(restricted_path);sequences=set(restricted.sequence.str.upper())
    old_subset=old[old.sequence.str.upper().isin(sequences)].reset_index(drop=True)
    new_subset=new[new.sequence.str.upper().isin(sequences)].reset_index(drop=True)
    assert restricted.fillna('').equals(old_subset.fillna(''))
    po=old_subset.groupby('sequence')[target].mean();pn=new_subset.groupby('sequence')[target].mean();positive_change=(pn-po).abs()>1e-6
    positive=pd.DataFrame({'old_mean_MIC':po,'new_mean_MIC':pn,'changed':positive_change});positive.index.name='sequence'
    positive[positive_change].to_csv(out/'external_127_positive_label_changes.csv',encoding='utf-8-sig')
    # EvoGradient's molar conversion and upper cap still leave changed targets.
    evo=ROOT.parents[1]/'baseline/sota_bench/retrain_evo_strict.py'
    import ast, math
    tree=ast.parse(evo.read_text());nodes=[node for node in tree.body if isinstance(node,ast.FunctionDef) and node.name=='mw']
    assignments=[node for node in tree.body if isinstance(node,ast.Assign) and any(isinstance(t,ast.Name) and t.id in ['AA_MW','MW','AA_WEIGHTS','RES_MW'] for t in node.targets)]
    # The independent check uses the exact literal function body and constants.
    scope={}
    for node in tree.body:
        if isinstance(node,ast.Assign):
            for target_node in node.targets:
                if isinstance(target_node,ast.Name):
                    try:scope[target_node.id]=ast.literal_eval(node.value)
                    except (ValueError,TypeError):pass
    exec(compile(ast.Module(body=nodes,type_ignores=[]),str(evo),'exec'),scope)
    mw=np.array([scope['mw'](s) for s in po.index]);evo_old=np.minimum(np.log10(po.to_numpy()*1000/mw),math.log10(8192)).astype(np.float32)
    evo_new=np.minimum(np.log10(pn.to_numpy()*1000/mw),math.log10(8192)).astype(np.float32)
    assert np.sum(evo_old!=evo_new)>0
    methods=[]
    for method,(filename,_,_) in METHODS.items():
        path=ROOT/'baseline/sota_bench'/filename
        if method=='EvoGradient':path=evo
        methods.append({'method':method,'conditions':'with_negatives;without_negatives','seeds_per_condition':10,'required_runs':20,
                        'training_driver':str(path),'sha256':sha(path),
                        'dependency':'shared restricted positive MIC targets; method-specific transformed/cached target must be regenerated'})
    methods.append({'method':'AMPLiT','conditions':'with_negatives;without_negatives','seeds_per_condition':10,'required_runs':20,
                    'training_driver':str(ROOT/'scripts/repair_v15/amplit.py'),'sha256':sha(ROOT/'scripts/repair_v15/amplit.py'),
                    'dependency':'training_data reads shared positive labels; label standardization and supervised weights must be refitted'})
    pd.DataFrame(methods).to_csv(out/'external_retraining_dependency.csv',index=False,encoding='utf-8-sig')
    evidence={'status':'READ_ONLY_SCOPE_CONFIRMED','timestamp':datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(timespec='seconds'),
              'raw_ACT_records_reproduced':399,'raw_ACT_Ecoli_labels_changed':int(((change[target+'_new']-change[target]).abs()>1e-6).sum()),
              'effective_TIGER_observations':7217,'effective_TIGER_labels_changed':int(diff.sum()),'TIGER_representatives_changed':int(representative.sum()),
              'folds':folds,'restricted_raw_rows':len(restricted),'restricted_raw_labels_changed':int(((new_subset['MIC_Escherichia_coli']-restricted['MIC_Escherichia_coli']).abs()>1e-6).sum()),
              'restricted_unique_sequences':len(po),'restricted_sequence_targets_changed':int(positive_change.sum()),
              'EvoGradient_changed_float32_capped_targets':int(np.sum(evo_old!=evo_new)),
              'required_TIGER_unique_runs':470,'required_external_two_condition_runs':420,'total_required_runs':890,
              'increase_over_original_660_runs':230,'scope_assumption':'Existing 40 ablation arms + 8 factorial cells + 1 gate-off, same seeds; fixed restricted 4929 sequence cohort; ACT mass alias only',
              'method_evidence_boundary':'positive-table numeric propagation, driver/cache dependency inspection, explicit Evo unit/cap check; no new per-method training or full-driver dry run',
              'production_or_queue_modified':False,'source_hashes':hashes}
    dump(out/'scope_evidence.json',evidence)
    blocks=[]
    def h(t):blocks.append(('h',t))
    def p(t):blocks.append(('p',t))
    def table(t):blocks.append(('t',t))
    h('ACT标签换算修复：重新训练范围核定')
    p('核查时间：'+evidence['timestamp']+'。只读复算原始399条ACT记录、现行去重和外部正样本聚合，不修改源表、训练配置、运行队列或模型。范围固定为当前登记的配置/种子与4,929条外部正序列；仅修正ACT质量别名，其他单位、端点、聚合规则保持既定口径。')
    p('结论：修复后需要470个TIGER种子级预测器和420个外部方法种子/条件预测器，共890个独立训练任务。相对原660个任务，登记范围增加230个：原计划复用的30个TIGER无p任务，以及此前未重训的20种外部方法无负条件200个任务。已完成或已启动且使用旧标签的受影响训练需替换；尚未启动的任务直接切换到冻结后的修复数据，不必先跑旧标签。')
    table([['任务类别','修复后独立任务数','原因'],['现有TIGER 40臂','40×10＝400','标签及部分去重代表改变，含三个无p模态臂'],['TIGER负样本8配置','新增60；共80个结果位置','P10、S11与已有臂共用20个结果，不重复训练'],['真正关闭相似度门','10','训练回归目标和配对标签改变'],['21种外部方法','21×2条件×10＝420','共同正样本标签改变，含负和无负都受影响；含AMPLiT'],['总计','470＋420＝890','TIGER每任务5折，共2,350个折模型身份；外部按各自既定配方']])
    h('一、标签传播的数值证据')
    table([['位置','复算结果'],['源表399条ACT记录','现有标签全部复现；165条E. coli源标签改变'],['固定原209条训练ACT观察','前次审计138条标签改变，尚未重新去重'],['重新运行现有去重规则','仍为7,217条观察；136条最终标签改变，另2条改选源记录'],['外部基准正样本原表','5,737条记录中129条MIC改变'],['按序列均值聚合后','4,929条正序列中127条目标改变'],['EvoGradient µM转换及上限截断',f"仍有{evidence['EvoGradient_changed_float32_capped_targets']}条float32回归目标变化，不能因单位不同而豁免"]])
    p('138与136不是矛盾：前者固定旧入选记录，后者重做最小MIC代表选择。两组原来标签相同的重复序列在ACT质量修复后改选未标注ACT的记录：RGRKGGRRKK从DBAASPS_8675切换为DBAASPS_8674；INWLKLGKKLLSAL从DBAASPS_2791切换为DBAASPS_2789。两组最终MIC保持旧值，但末端字段和理化输入改变。该复算沿用现有sequence×CFU去重定义，并未裁定两种化学末端是否应作为独立分子实体；若进一步改变该定义，需要另行核定样本集合。')
    table([['折','训练集改变的标签数','验证集改变的标签数']]+[[str(r['fold']),str(r['changed_train_labels']),str(r['changed_validation_labels'])] for r in folds])
    h('二、为什么无p模态也必须重训')
    p('上一次ACT修复只改变输入向量的charge/charge_density，无p分支的序列、结构模型可经验证后复用。此次改变的是监督目标MIC及其log2差值；所有模态的回归损失、验证选模和校准都依赖目标值。因此s4_mod_s、s4_mod_g、s4_mod_sg三个配置的30个种子不能继续用“未使用p”作为免重训依据。需要从规范起点重新训练受影响监督阶段，不以旧回归权重追加几个epoch冒充同配方重训。')
    p('TIGER的训练配对分箱、配额、差值标签及验证/校准输入应使用新标签重建；图缓存中的目标值也必须更新。相同序列/CFU身份保留使原GroupKFold分组可继续使用，但完整输入和折内标准化需重新核验。两个去重代表的末端变化也意味着含p分类预训练输入可能变化，不能统一承诺复用所有N2预训练权重。')
    h('三、为什么外部方法的两条件都必须重训')
    p('此次错误在正样本MIC，含负与无负条件均使用这些正标签。此前负池冲突修复只影响含负条件，这个豁免现在不再适用。21种外部方法：'+ '、'.join(r['method'] for r in methods)+'。AMPLiT已在这21种中，不额外重复计数；TIGER-single计入TIGER部分。')
    p('方法依赖依据为共享正样本表数值复算与实际训练入口/缓存读取代码，另对EvoGradient执行精确单位转换/截断后的目标差异检查。尚未重新运行全部方法驱动或逐种子fit截停；正式执行前仍应逐方法验证最终目标数组、标签标准化器和缓存身份，避免某些缓存继续带旧标签。')
    h('四、可以保留的工作及后续顺序')
    p('可复用已验证不依赖MIC的原始PDB、序列token、冻结PLM嵌入、纯结构和序列描述符，以及保持不变的划分身份。缓存若混存标签，必须更新目标并重新登记哈希。训练子集相同、无标签拟合的固定Word2Vec/TF-IDF等表征可经身份核验后复用；监督训练权重、监督特征选择、目标标准化器和校准器按实际依赖重新拟合。已完成的GPU环境修复和数值一致性验证仍然有用。')
    p('建议先冻结标签修复与去重口径，完成原始JSON→训练源表→有效池→外部均值标签→逐方法实际目标数组的验收，再更新任务注册表。然后安排TIGER 470个及外部420个任务，最后重做外测四指标、配对统计、校准输出验收和统一推理计时。性能计时与图文生成是后续评估工作，不另计训练任务。')
    p('890是修复后所需任务总规模，不能与当前已完成任务数相加。每种子成本依方法而异，不能用TIGER平均训练时间乘890估计工期。该范围核定不自动启动新训练，也不对原运行队列执行取消或重启。')
    p('复算代码：scripts/repair_v15/audit_act_retraining_scope.py。附表包括TIGER_effective_changes.csv、TIGER_two_representative_changes.csv、TIGER_fold_label_changes.csv、external_127_positive_label_changes.csv、external_retraining_dependency.csv；scope_evidence.json保存范围假设与源文件哈希。')
    doc=Document();style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10);style.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    md=[];texts=[]
    for kind,val in blocks:
        if kind=='h':doc.add_heading(val,level=1);md.extend(['# '+val,'']);texts.append(val)
        elif kind=='p':doc.add_paragraph(val);md.extend([val,'']);texts.append(val)
        else:
            t=doc.add_table(rows=0,cols=len(val[0]));t.style='Table Grid'
            for row in val:
                for cell,text in zip(t.add_row().cells,row):cell.text=text;texts.append(text)
            md.extend(['| '+' | '.join(val[0])+' |','| '+' | '.join(['---']*len(val[0]))+' |'])
            md.extend('| '+' | '.join(row)+' |' for row in val[1:]);md.append('')
    base=out/'ACT_retraining_scope_ZH';base.with_suffix('.md').write_text('\n'.join(md));doc.save(base.with_suffix('.docx'))
    check=Document(base.with_suffix('.docx'));content='\n'.join([p.text for p in check.paragraphs]+[c.text for t in check.tables for row in t.rows for c in row.cells])
    assert all(text in content for text in texts)
    assert all(sha(path)==value for path,value in hashes.items())
    print({k:v for k,v in evidence.items() if k!='source_hashes'});print(out)


if __name__=='__main__':main()
