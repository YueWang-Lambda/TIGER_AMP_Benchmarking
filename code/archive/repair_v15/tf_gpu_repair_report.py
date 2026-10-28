"""Check CPU/GPU endpoint parity and emit the same Markdown/Word report."""
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo
import statistics
import numpy as np
import pandas as pd
from common import ROOT, WORK, REPORT, load, dump
from external_metrics import from_peptides, evaluate
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn


def read_predictions(path):
    frame = pd.read_csv(path)
    assert len(frame) == 208 and frame.sequence.is_unique
    series = frame.set_index('sequence').log10mic
    assert np.isfinite(series).all()
    return series


def main():
    out = WORK/'tf_gpu_repair'
    results = {}
    for method, seed in [('AMPlify', 120), ('MBC-Attention', 123)]:
        probes = {dev: load(out/f'{method}_{dev}.json') for dev in ['cpu', 'gpu']}
        assert all(r['status'] == 'PASS' for r in probes.values())
        assert probes['cpu']['checkpoint_sha256'] == probes['gpu']['checkpoint_sha256']
        preds = {dev: read_predictions(out/f'{method}_{dev}_predictions.csv') for dev in probes}
        archived = read_predictions(WORK/f'external_runs/{method}/seed{seed}/predictions.csv')
        assert all(set(p.index) == set(archived.index) for p in preds.values())
        pairs = {dev: from_peptides(p, unit='log10_ug_mL') for dev, p in preds.items()}
        mets = {dev: evaluate(p)[0] for dev, p in pairs.items()}
        peptide_err = float(np.max(np.abs(preds['gpu'].reindex(preds['cpu'].index)-preds['cpu'])))
        pair_err = float(np.max(np.abs(pairs['gpu'].y_pred-pairs['cpu'].y_pred)))
        archives = {dev: float(np.max(np.abs(p.reindex(archived.index)-archived))) for dev, p in preds.items()}
        diff = {m: abs(mets['gpu']['macro'][m]-mets['cpu']['macro'][m]) for m in ['PCC', 'RSE']}
        # Absolute tolerances specified for float32 kernel parity; not an accuracy
        # acceptance criterion and not evidence of identical training trajectories.
        thresholds = {'peptide_log10mic': 1e-4, 'pair_log2_delta': 7e-4,
            'macro_PCC': 1e-4, 'macro_RSE': 1e-4, 'archived_peptide_log10mic': 1e-4}
        assert peptide_err <= thresholds['peptide_log10mic']
        assert pair_err <= thresholds['pair_log2_delta']
        assert all(v <= thresholds['archived_peptide_log10mic'] for v in archives.values())
        assert all(diff[m] <= thresholds['macro_'+m] for m in diff)
        assert all(r['tf32_enabled'] is False and r['graph_options']['arithmetic_optimization'] is False for r in probes.values())
        if method == 'AMPlify':
            assert probes['gpu']['functions_run_eagerly'] is True
            assert load(out/'amplify_stability_eager.json')['pass'] is True
        assert probes['cpu']['gradient_variable_names'] == probes['gpu']['gradient_variable_names']
        gradients = {dev: np.load(out/f'{method}_{dev}_gradients.npz') for dev in probes}
        assert gradients['cpu'].files == gradients['gpu'].files
        max_grad = 0.; squared_diff = 0.; squared_cpu = 0.
        for key in gradients['cpu'].files:
            cpu = gradients['cpu'][key].astype(float); gpu = gradients['gpu'][key].astype(float)
            assert cpu.shape == gpu.shape
            max_grad = max(max_grad, float(np.max(np.abs(cpu-gpu))))
            squared_diff += float(np.sum((cpu-gpu)**2)); squared_cpu += float(np.sum(cpu**2))
        grad_rel = float(np.sqrt(squared_diff/max(squared_cpu, 1e-30)))
        assert max_grad <= 1e-4 and grad_rel <= 1e-3
        results[method] = {'status': 'PASS', 'seed': seed, 'parameters': probes['gpu']['parameters'],
            'n_peptides': 208, 'n_pairs': 2211, 'n_families': 11,
            'max_peptide_log10mic_abs_diff': peptide_err,
            'max_pair_log2_abs_diff': pair_err, 'macro_abs_diff': diff,
            'max_abs_diff_vs_formal_archived_peptides': archives,
            'metrics': mets, 'absolute_tolerances': thresholds,
            'deterministic_gradient_max_abs_diff': max_grad,
            'deterministic_gradient_relative_l2_diff': grad_rel,
            'gpu_execution': 'eager' if method == 'AMPlify' else 'graph',
            'cpu_batch_median_seconds': probes['cpu']['train_batch_median_seconds'],
            'gpu_batch_median_seconds': probes['gpu']['train_batch_median_seconds'],
            'observed_batch_speed_ratio': probes['cpu']['train_batch_median_seconds']/probes['gpu']['train_batch_median_seconds'],
            'gpu_tensor_memory_peak_mib': probes['gpu']['tf_memory_bytes']['peak']/1024**2}
    stamp = datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(timespec='seconds')
    result = {'status': 'PASS', 'captured_at': stamp, 'methods': results,
        'campaign_queue_modified': False, 'active_training_interrupted': False,
        'no_formal_predictions_or_cost_timing_overwritten': True}
    dump(out/'validation_summary.json', result)
    dump(out/'RUNTIME_READY.json', {'status': 'VALIDATED_AVAILABLE_NOT_IN_FORMAL_QUEUE',
        'prefix': '/home/ubuntu/anaconda3/envs/mbc_att_gpu_v15',
        'launcher': str(ROOT/'scripts/repair_v15/tf_gpu_launcher.py'),
        'mandatory_gpu_execution': {'AMPlify': 'eager', 'MBC-Attention': 'graph'},
        'tf32': False, 'arithmetic_optimization': False,
        'tensorflow_allocation_cap_mib': 2048,
        'proposed_full_process_reservation_mib': 3072,
        'validation_summary': str(out/'validation_summary.json'),
        'note': 'Use the launcher; activating the clone alone does not fix AMPlify GPU graph parity.'})
    active_profile=load(WORK/'tf_gpu_runtime.json') if (WORK/'tf_gpu_runtime.json').exists() else None
    if active_profile:
        ready=load(out/'RUNTIME_READY.json')
        ready.update(status='ACTIVE_GPU_QUEUE',tensorflow_allocation_cap_mib=active_profile['memory_cap_mib'],
            proposed_full_process_reservation_mib=active_profile['reservation_mib'],
            active_profile=str(WORK/'tf_gpu_runtime.json'),
            migration_verification=str(Path(active_profile['migration_directory'])/'handoff_verification.json'))
        dump(out/'RUNTIME_READY.json',ready)
    blocks = []
    def h(s): blocks.append(('h', s))
    def p(s): blocks.append(('p', s))
    def t(rows): blocks.append(('t', rows))
    h('AMPlify与MBC-Attention的TensorFlow GPU加载修复')
    p('报告更新时间：'+stamp+'。两种方法均已通过真实GPU前向、训练更新及CPU/GPU值级一致性检查。兼容性诊断阶段未修改原训练环境或正式结果；后续用户授权的GPU迁移见本报告末节及独立迁移报告。')
    h('根因与修复')
    p('原mbc_att环境已安装CUDA 11.8与cuDNN 8.6，而调用绝对路径的Python不会执行conda activate钩子。原启动环境的LD_LIBRARY_PATH只有/usr/local/cuda/lib64:，未包含mbc_att/lib及其NVIDIA wheel库目录。TensorFlow 2.12.1的动态加载路径因此找不到部分GPU依赖，最终注册GPU失败。直接用ctypes加载个别库可以成功，并不能证明TensorFlow自身的动态加载器也能解析全部依赖。')
    p('对照实验证实：同一原环境、同一TensorFlow，显式加入mbc_att/lib后，从GPU列表为空变为成功识别RTX 4090 D。修复不需要升级TensorFlow、Keras或NVIDIA驱动。当前正式训练包装器还显式传递--cpu并设置CUDA_VISIBLE_DEVICES=-1；仅解决库路径不会自动改变这些已启动训练的执行设备。')
    p('已将mbc_att克隆为独立环境/home/ubuntu/anaconda3/envs/mbc_att_gpu_v15，修正克隆环境中17个CUDA库软链接，使其指向克隆自身。专用启动器明确设置CONDA_PREFIX、PATH、CUDA_VISIBLE_DEVICES及CUDA 11库路径，避免继承其他环境的CUDA 12库路径；保持TensorFlow 2.12.1、Keras 2.12.0、CUDA runtime 11.8.89和cuDNN 8.6.0.163。')
    h('值级检查拦下的GPU兼容性问题')
    p('仅恢复库加载还不足以验收。首轮默认GPU运行中，AMPlify同权重逐肽log10MIC最大差为0.8394389，宏PCC从CPU的0.542198变为0.520696；该结果未被接受，也未写入正式实验。层级追踪发现差异位于带掩码的反向LSTM分支；正向LSTM仅有约3e-7差异。独立ReverseSequence算子检查通过，eager模式双向LSTM也与CPU在3e-7以内一致。')
    p('关闭TF32后，GPU图模式仍出现约0.836的log10MIC差异。图优化对照显示：关闭layout优化无效；仅关闭arithmetic_optimization曾将最大差降至2.15e-6，但新进程完整复验仍有9条肽出现差异，最大0.1995964。因此“仅关闭算术优化”未获验收。进一步关闭整个meta optimizer，或改用eager执行，均通过原顺序、重复、逆序、随机重排的4次完整208肽检查。')
    p('最终AMPlify GPU配置采用eager执行（tf.config.run_functions_eagerly(True)），绕开存在差异的图执行路径，同时保留cuDNN；MBC-Attention采用已验收的graph执行。启动器默认eager，并阻止已识别的AMPlify正式训练/诊断命令选择graph。没有改写模型层、权重、掩码公式、损失或训练超参数。这是本机TF 2.12.1组合的兼容性处理，尚未将异常定位到框架底层某一条重写规则，不能将其概括为所有GPU或所有TensorFlow版本均有该问题。')
    p('启动器同时通过TensorFlow开关及NVIDIA_TF32_OVERRIDE=0禁用TF32，以FP32执行，减少与CPU参考的精度差异。两种模型均按最终配置重新验证。默认GPU失败输出与优化开关对照保存在initial_default_gpu/及graph_options_probe.json，可复核修复前后差异；不能把“能运行GPU”替代为“数值正确”。')
    h('真实模型验证')
    p('AMPlify加载本轮已完成的seed120权重；MBC-Attention加载本轮已完成的seed123权重。每个模型在CPU和GPU上读取相同权重、相同208条外测肽，按规范边清单计算2,211条log2差分并核对11家族宏PCC/RSE，同时对照正式归档的逐肽预测。模型架构、掩码、标签逆变换与batch size 32均保持不变。')
    rows = [['方法', '208肽log10MIC最大差', '2,211边log2差分最大差', '宏PCC差', '宏RSE差']]
    for method, r in results.items():
        rows.append([method, f"{r['max_peptide_log10mic_abs_diff']:.3e}",
            f"{r['max_pair_log2_abs_diff']:.3e}", f"{r['macro_abs_diff']['PCC']:.3e}", f"{r['macro_abs_diff']['RSE']:.3e}"])
    t(rows)
    p('上述差值为同一权重在CPU/GPU上的绝对差。预设验收容差：逐肽log10MIC≤1e-4、逐边log2差分≤7e-4、宏PCC和RSE差各≤1e-4；两方法均通过，且均与本轮正式逐肽归档在1e-4内一致。验证脚本严格核对完整身份集合，未删除种子、家族、肽或边。')
    p('训练检查独立抽取256条训练语料（4,929正样本＋7,949去冲突负样本池），运行3轮、每轮8个batch。外部测试集只用于前向一致性检查，不参与训练更新。两个模型在两设备上均损失有限、可训练权重确实更新；GPU可训练变量位于GPU:0，实际加载的cuDNN/cuBLAS库来自克隆环境。诊断权重只存在内存中，未保存为正式模型；前后权重文件SHA256一致。')
    rows = [['方法', '全部参数梯度最大绝对差', '梯度差L2／CPU梯度L2']]
    for method, r in results.items():
        rows.append([method, f"{r['deterministic_gradient_max_abs_diff']:.3e}", f"{r['deterministic_gradient_relative_l2_diff']:.3e}"])
    t(rows)
    p('另在同一初始权重、同一32条训练样本上，以training=False关闭随机dropout做一次确定性反向传播检查，按变量名逐项核对全部梯度。CPU参考和MBC-Attention GPU为graph执行，AMPlify GPU为最终的eager执行。最大绝对差≤1e-4、相对L2差≤1e-3均通过。随后3轮model.fit仍使用原配方的训练dropout；该确定性诊断不修改正式训练方式。')
    h('加速与显存观察')
    rows = [['方法', 'CPU训练batch中位数', 'GPU训练batch中位数', '本次耗时比CPU/GPU', 'TF张量分配峰值']]
    for method, r in results.items():
        rows.append([method, f"{r['cpu_batch_median_seconds']:.4f}秒", f"{r['gpu_batch_median_seconds']:.4f}秒",
            f"{r['observed_batch_speed_ratio']:.2f}", f"{r['gpu_tensor_memory_peak_mib']:.1f} MiB"])
    t(rows)
    p('统计排除第一轮编译预热，来自后两轮16个batch。诊断与现有训练共享CPU/GPU，不能作为论文推理计时、完整训练加速倍数或完成时间保证。GPU与CPU的浮点运算顺序、随机数及cuDNN实现存在差异；前向一致性不等于重训权重逐位一致，也不证明两设备训练轨迹完全相同。')
    p('本次GPU诊断按2,048 MiB TensorFlow分配上限运行；驱动/CUDA上下文有额外开销，nvidia-smi观察到AMPlify进程约2,454 MiB，不能把TensorFlow分配上限理解成整进程显存上限。两方法依次使用同一张卡的第5个进程位置；调度器识别了诊断进程，正式队列仍受每卡5进程及显存预算约束。启动器在执行前检查进程数和可用显存，并在GPU不可见时直接失败，避免静默退回CPU。')
    h('当前工作安排与复现入口')
    if active_profile:
        p('用户随后明确授权取消未完成CPU任务并切换GPU。正式队列现使用独立GPU环境：3个AMPlify CPU尝试已终止归档，8个已完成CPU结果保留，其余12个任务转入GPU队列。迁移时进一步验证了每任务1 GiB TensorFlow分配上限，按整进程2 GiB调度预留；详细进度与运行核验见TensorFlow_GPU_migration_ZH.md/.docx。本报告上面的速度和显存表保留最初2 GiB诊断记录，不充当正式推理计时。')
    else:
        p('当前正在运行的CPU训练继续完成，正式队列未自动切换设备。独立环境与启动器已经可用。后续接入正式调度时，应去掉目标训练命令的--cpu，按整进程暂预留3 GiB并登记实际训练设备；必须使用启动器的已验收配置。只激活新环境或更换Python路径，仍不足以修复AMPlify图模式差异。完整推理计时仍应在空闲设备下按统一终点另行执行。')
    p('GPU复现命令（项目根目录执行；已有5个GPU进程时启动器会拒绝）：/home/ubuntu/anaconda3/envs/ccseg/bin/python scripts/repair_v15/tf_gpu_launcher.py --gpu 0 --execution eager --memory-mib 2048 --target scripts/repair_v15/validate_tf_gpu.py -- --method AMPlify --device gpu。MBC-Attention使用--execution graph并替换方法名；CPU参考使用--gpu -1、--execution graph和--device cpu。')
    t([['内容', '相对项目根目录的路径'],
        ['独立环境启动器', 'scripts/repair_v15/tf_gpu_launcher.py'],
        ['真实模型诊断', 'scripts/repair_v15/validate_tf_gpu.py'],
        ['终点核对及报告生成', 'scripts/repair_v15/tf_gpu_repair_report.py'],
        ['根因对照', 'artifacts/repair_v15/environment_audit/mbc_att.json；mbc_att_explicit_libpath.json'],
        ['克隆环境修复记录', 'artifacts/repair_v15/tf_gpu_repair/environment_fix.json'],
        ['验收总表', 'artifacts/repair_v15/tf_gpu_repair/validation_summary.json'],
        ['GPU默认设置失败对照', 'artifacts/repair_v15/tf_gpu_repair/initial_default_gpu/failed_parity_summary.json'],
        ['层级追踪及图优化对照', 'scripts/repair_v15/trace_amplify_gpu_layers.py；probe_tf_reverse.py；probe_tf_graph_options.py'],
        ['逐设备记录、日志及预测', 'artifacts/repair_v15/tf_gpu_repair/{AMPlify,MBC-Attention}_{cpu,gpu}.{json,log}及*_predictions.csv']])
    doc = Document()
    normal = doc.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    md = []; expected = []
    for kind, content in blocks:
        if kind == 'h':
            doc.add_heading(content, level=1); md.extend(['# '+content, '']); expected.append(content)
        elif kind == 'p':
            doc.add_paragraph(content); md.extend([content, '']); expected.append(content)
        else:
            table = doc.add_table(rows=0, cols=len(content[0])); table.style = 'Table Grid'
            for row in content:
                for cell, text in zip(table.add_row().cells, row): cell.text = text; expected.append(text)
            md.extend(['| '+' | '.join(content[0])+' |', '| '+' | '.join(['---']*len(content[0]))+' |'])
            md.extend('| '+' | '.join(row)+' |' for row in content[1:]); md.append('')
    dest = REPORT/'TensorFlow_GPU_repair_ZH'
    dest.with_suffix('.md').write_text('\n'.join(md)); doc.save(dest.with_suffix('.docx'))
    check = Document(dest.with_suffix('.docx'))
    actual = '\n'.join([p.text for p in check.paragraphs]+[c.text for t in check.tables for row in t.rows for c in row.cells])
    assert all(s in actual for s in expected)
    print(result)
    print(dest)


if __name__ == '__main__':
    main()
