"""Reconcile model environments with launch manifests and owned live processes."""
import ast
import csv
from datetime import datetime
from zoneinfo import ZoneInfo
from common import *
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn


def main():
    home=WORK/'environment_audit'
    tree=ast.parse((ROOT/'scripts/repair_v15/external_one.py').read_text())
    methods=next(ast.literal_eval(n.value) for n in tree.body if isinstance(n,ast.Assign)
                 and any(isinstance(t,ast.Name) and t.id=='METHODS' for t in n.targets))
    migrated=(WORK/'tf_gpu_runtime.json').exists()
    env_names=['ccseg','mbc_att','ampacti','ania']+(['mbc_att_gpu_v15'] if migrated else [])
    probes={name:load(home/f'{name}.json') for name in env_names}
    preserved={}
    if migrated:
        directory=Path(load(WORK/'tf_gpu_runtime.json')['migration_directory'])
        preserved=load(directory/'completed_preserved_hashes.json')
    live=load(home/'live_processes.json'); rows=[]; mismatches=[]
    for method,(script,kind,env) in methods.items():
        expected=f'/home/ubuntu/anaconda3/envs/{env}/bin/python'
        manifests=[]
        for path in sorted((WORK/'external_runs'/method).glob('seed*/invocation.json')):
            data=load(path)
            ok=data['cmd'][0]==expected
            device=data['env'].get('CUDA_VISIBLE_DEVICES')
            device_ok=(device=='-1') if kind in ['cpu','tf'] else device in ['0','1']
            historical_cpu=(str(path.parent.relative_to(ROOT)) in preserved and
                (path.parent/'DONE.json').exists() and load(path.parent/'DONE.json')['status']=='PASS')
            if historical_cpu:
                ok=data['cmd'][0]=='/home/ubuntu/anaconda3/envs/mbc_att/bin/python'
                device_ok=device=='-1'
            if kind=='tf_gpu' and not historical_cpu and (path.parent/'runtime_environment.json').exists():
                actual=load(path.parent/'runtime_environment.json')
                ok=ok and actual['prefix']=='/home/ubuntu/anaconda3/envs/mbc_att_gpu_v15'
                device_ok=device_ok and actual['physical_gpu_index']==int(device)
                device_ok=device_ok and actual['execution']==('eager' if method=='AMPlify' else 'graph')
            rec={'seed':path.parent.name,'python':data['cmd'][0],'cuda_visible_devices':device,
                 'preserved_completed_cpu_seed':historical_cpu,
                 'interpreter_match':ok,'device_selection_match':device_ok,'manifest':str(path.relative_to(ROOT))}
            manifests.append(rec)
            if not ok or not device_ok:mismatches.append({'method':method,**rec})
        model_live=[x for x in live if x['job_id'].startswith(method+'_with_negatives_')
                    and x.get('driver','').endswith(script)]
        for rec in model_live:
            if rec['executable']!=str(Path(expected).resolve()):mismatches.append({'method':method,'live':rec})
        rows.append({'method':method,'environment':env,'python':probes[env]['python'],
                     'device':'GPU（保留已完成CPU种子）' if kind=='tf_gpu' else 'CPU' if kind in ['cpu','tf'] else 'GPU',
                     'interpreter':expected,'driver':script,'launch_records':len(manifests),
                     'observed_invocations':manifests,'live_driver_pids':[x['pid'] for x in model_live],
                     'evidence':'启动记录核对'+('；实际进程核对' if model_live else '')})
    amplit_live=[x for x in live if x['job_id'].startswith('AMPLiT_')
                 and str(x.get('driver','')).endswith('scripts/repair_v15/amplit.py')]
    for rec in amplit_live:
        if rec['executable']!=probes['ccseg']['real_executable']:mismatches.append({'method':'AMPLiT','live':rec})
    rows.append({'method':'AMPLiT','environment':'ccseg','python':probes['ccseg']['python'],'device':'GPU',
                 'interpreter':'/home/ubuntu/anaconda3/envs/ccseg/bin/python',
                 'driver':'scripts/repair_v15/amplit.py','launch_records':None,
                 'live_driver_pids':[x['pid'] for x in amplit_live],
                 'evidence':'execute_job/common固定解释器；实际进程核对' if amplit_live else 'execute_job/common固定解释器'})
    assert len(rows)==21 and all(r['launch_records'] for r in rows if r['method']!='AMPLiT')
    assert all(p['framework_import']['status']=='PASS' for p in probes.values())
    tag_mismatches=[]
    for rec in live:
        prefix=str(Path(rec['executable']).parent.parent)
        tag=rec['env'].get('CONDA_PREFIX')
        if tag and tag!=prefix:tag_mismatches.append({'job_id':rec['job_id'],'pid':rec['pid'],
                'actual_interpreter_prefix':prefix,'inherited_CONDA_PREFIX':tag})
    stamp=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(timespec='seconds')
    snapshot={'captured_at':stamp,'scope':'v15 training drivers and in-driver predictions; canonical metrics separately in ccseg',
              'method_count':21,'launch_record_count':sum(r['launch_records'] or 0 for r in rows),
              'interpreter_mismatches':mismatches,'inherited_environment_tag_mismatches':tag_mismatches,
              'methods':rows,'environments':probes,
              'status':'INTERPRETERS_MATCH_WITH_NOTES' if not mismatches else 'MISMATCH'}
    dump(home/'environment_audit.json',snapshot)
    columns=['method','environment','python','device','interpreter','driver','launch_records','evidence']
    with (home/'method_environment_table.csv').open('w',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=columns);writer.writeheader()
        writer.writerows({k:r[k] for k in columns} for r in rows)
    blocks=[]
    def h(x):blocks.append(('h',x))
    def p(x):blocks.append(('p',x))
    def table(x):blocks.append(('t',x))
    h('外部基准测试conda环境核查')
    p('核查时间：'+stamp+'。核查范围为本轮v15的21种外部方法模型训练及驱动内预测，包含AMPLiT的两个训练条件。历史复用结果当时的完整环境不能仅由当前安装状态反推。')
    p(f"已核对20种外部方法的{snapshot['launch_record_count']}份实际启动记录，另核对AMPLiT的固定启动链与在运行的真实Python进程。{len(probes)}个环境均完成解释器、sys.prefix、包版本和主要框架的只读导入检查；训练解释器/设备选择不匹配共{len(mismatches)}项。本次核查只读取证据；此前用户授权的TensorFlow GPU迁移已按逐种子记录纳入。")
    table([['方法','conda环境','训练/驱动内预测设备','核查依据']]+
          [[r['method'],r['environment'],r['device'],r['evidence']] for r in rows])
    h('环境版本与真实路径')
    environment_rows=[['环境','Python','主要框架','GPU能力探测']]
    for name in ['ccseg','ania','mbc_att','ampacti']+(['mbc_att_gpu_v15'] if migrated else []):
        d=probes[name]
        if name in ['ccseg','ania']:
            framework='PyTorch '+d['framework_import']['version'];gpu='CUDA '+d['cuda']['compiled_version']+'；可用'
        elif name=='mbc_att':framework='TensorFlow '+d['versions']['tensorflow']+' / Keras '+d['versions']['keras'];gpu='历史未带激活库路径的探测未识别GPU'
        elif name=='mbc_att_gpu_v15':framework='TensorFlow '+d['versions']['tensorflow']+' / Keras '+d['versions']['keras'];gpu='GPU加载与两模型数值/梯度验收通过'
        else:framework='deep-forest '+d['versions']['deep-forest'];gpu='树模型使用CPU'
        environment_rows.append([name,d['python'],framework,gpu])
    table(environment_rows)
    table([['环境','解释器路径']]+[[name,probes[name]['executable']] for name in probes])
    p('ccseg共承担17种外部方法，其中EIPpred使用CPU，其他16种使用GPU；ANIA另用ania环境和GPU；AMPActiPred用ampacti和CPU。'+('MBC-Attention和AMPlify的后续任务使用mbc_att_gpu_v15与GPU，分别保留6个和2个已完成CPU种子；合计后续19种GPU方法、2种CPU方法。' if migrated else 'MBC-Attention、AMPlify用mbc_att和CPU；合计17种GPU方法、4种CPU方法。')+'TIGER两臂另由ccseg运行，不计入21种外部方法。')
    h('发现一：原启动方式无法加载TensorFlow的GPU依赖')
    p('mbc_att安装的TensorFlow 2.12.1本身是CUDA构建，构建信息为CUDA 11.8/cuDNN 8。独立探测明确设置CUDA_VISIBLE_DEVICES=0且未设置--cpu，仍返回GPU列表为空，并输出“Cannot dlopen some GPU libraries”和“Skipping registering GPU devices”。所以问题不是机器没有GPU，也不能仅归因于训练命令的--cpu；当前环境确实无法加载所需GPU运行库。')
    p('迁移前，AMPlify和MBC-Attention的正式任务显式使用--cpu及CUDA_VISIBLE_DEVICES=-1，与此前计时披露一致。已完成CPU种子的启动记录继续保留；后续GPU任务以新启动器及runtime_environment.json记录实际环境和执行方式。')
    if (WORK/'tf_gpu_repair/validation_summary.json').exists():
        p('2026-10-12后续修复更正：CUDA 11.8/cuDNN 8.6软件包原本已经安装，直接调用Python绕过conda激活钩子，缺少库搜索路径才是无法加载的原因；此前“依赖不完整”不能理解为尚未安装。独立克隆环境mbc_att_gpu_v15已通过实际GPU验证，AMPlify必须使用eager执行避开图模式数值差异，MBC-Attention可使用已验收的graph执行。'+('用户随后授权迁移：3个未完成CPU尝试取消归档，8个已完成CPU结果保留；剩余12个任务转入GPU队列。' if migrated else '该环境尚未接入正式训练队列，原任务继续使用CPU。')+'完整配置与证据见同目录TensorFlow_GPU_repair_ZH.md/.docx和TensorFlow_GPU_migration_ZH.md/.docx。')
    h('发现二：继承的conda标签不能用于判定实际解释器')
    p(f'在本次存活进程快照中发现{len(tag_mismatches)}个进程的CONDA_PREFIX标签与实际解释器前缀不同。标签不匹配的任务为：'+('、'.join(sorted({r['job_id'] for r in tag_mismatches})) if tag_mismatches else '无')+'。非TensorFlow启动器仍可能直接调用其他环境的Python而继承父进程标签；实际解释器与标签应分别记录。新TensorFlow GPU启动器同时设置正确的CONDA_PREFIX、PATH和库搜索路径。')
    p('迁移前mbc_att解释器的独立探测sys.prefix和TensorFlow模块路径均指向mbc_att，旧AMPlify日志也显示该环境下的Keras路径；迁移后模型解释器和runtime_environment.json指向mbc_att_gpu_v15。外层调度与统一指标评估继续使用ccseg。应以真实解释器、sys.prefix及模块路径登记环境；仅查看终端提示符、CONDA_PREFIX或外层调度进程会误判。')
    h('分阶段环境与命名说明')
    p('模型训练及驱动内部逐肽预测使用上述各自环境；启动器和统一外测差分、四指标计算使用ccseg。这是有意的分阶段执行，AMPActiPred/ANIA/AMPlify的外层包装进程显示ccseg并不代表模型训练用错环境。')
    p('本机还存在hydramp、deepAMP、AMPDesigner等命名环境，但当前登记的HydrAMP回归适配、DeepAMP PyTorch移植及AMP-Designer训练入口使用ccseg。是否使用同名环境不是正确性的判据，应以实际代码依赖、已登记实现和启动证据为准。AMPLiT（PyTorch）与AMPlify（TensorFlow）是不同方法，不能混淆。')
    p('ANIA的专用R/kaos依赖位于ania环境，其驱动明确设置R_HOME；当前训练使用已构建的特征缓存。特征缓存的历史生成环境与本轮训练环境是不同的溯源层次，本报告不声称已为所有历史缓存重建当时的完整软件环境。')
    h('证据与后续记录')
    table([['内容','相对项目根目录的路径'],['每方法环境表','artifacts/repair_v15/environment_audit/method_environment_table.csv'],['完整核查JSON','artifacts/repair_v15/environment_audit/environment_audit.json'],['四环境解释器/版本探测','artifacts/repair_v15/environment_audit/{ccseg,ania,mbc_att,ampacti}.json'],['存活工作进程快照','artifacts/repair_v15/environment_audit/live_processes.json'],['逐种子启动记录','artifacts/repair_v15/external_runs/<方法>/seed<种子>/invocation.json'],['环境探测脚本','scripts/repair_v15/probe_conda_environment.py'],['核查及报告生成','scripts/repair_v15/audit_conda_environments.py']])
    p('建议在v15复现包中同时保留目标解释器、实际sys.prefix、核心包版本、CUDA/cuDNN构建与可用设备信息，并分别标注训练、特征生成及计时环境。当前环境版本快照已保存；它不能替代未来新环境的依赖锁文件或历史任务当时的安装快照。')
    doc=Document();style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    md=[];expected=[]
    for kind,value in blocks:
        if kind=='h':doc.add_heading(value,level=1);md.extend(['# '+value,'']);expected.append(value)
        elif kind=='p':doc.add_paragraph(value);md.extend([value,'']);expected.append(value)
        else:
            t=doc.add_table(rows=0,cols=len(value[0]));t.style='Table Grid'
            for row in value:
                for c,txt in zip(t.add_row().cells,row):c.text=txt;expected.append(txt)
            md.extend(['| '+' | '.join(value[0])+' |','| '+' | '.join(['---']*len(value[0]))+' |'])
            md.extend('| '+' | '.join(row)+' |' for row in value[1:]);md.append('')
    base=REPORT/'benchmark_conda_environment_audit_ZH'
    base.with_suffix('.md').write_text('\n'.join(md));doc.save(base.with_suffix('.docx'))
    actual=Document(base.with_suffix('.docx'))
    text='\n'.join([p.text for p in actual.paragraphs]+[c.text for t in actual.tables for row in t.rows for c in row.cells])
    assert all(v in text for v in expected)
    print(json.dumps({'report':str(base),'methods':21,'launch_records':snapshot['launch_record_count'],
        'interpreter_mismatches':len(mismatches),'tag_mismatches':len(tag_mismatches)},ensure_ascii=False))


if __name__=='__main__':main()
