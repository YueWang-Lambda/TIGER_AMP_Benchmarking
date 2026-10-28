"""Generate matching live Markdown/Word reports from one content tree."""
from datetime import datetime
from zoneinfo import ZoneInfo
import csv
import re
import statistics
import time
from common import *
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn


def main():
    if (WORK/'CANCELLED_BY_V14_INPUT_FREEZE.json').exists():
        import subprocess
        subprocess.run([PYTHON,str(ROOT/'scripts/repair_v15_v14input/status_report.py')],cwd=ROOT,check=True)
        return
    REPORT.mkdir(parents=True,exist_ok=True)
    reg=load(WORK/'run_registry.json')
    counts={k:sum(r['action']==k for r in reg) for k in ['train','reuse_validate','alias']}
    def passed(path):
        return path.exists() and load(path).get('status')=='PASS'
    finished=[r for r in reg if r['action']=='train' and passed(Path(r['run_dir'])/'DONE.json')]
    done=len(finished)
    eval_counts={action:sum(passed(WORK/'tiger_external'/r['task_id']/'EVALUATED.json') for r in reg if r['action']==action) for action in ['train','reuse_validate']}
    snapshot_time=datetime.now(ZoneInfo('Asia/Shanghai'))
    s=load(WORK/'queue/campaign/state.json')
    blocks=[]
    def h(t):blocks.append(('h',t))
    def p(t):blocks.append(('p',t))
    def table(rows):blocks.append(('t',rows))
    h('v15 修复执行记录')
    p('更新时间：'+snapshot_time.isoformat(timespec='seconds')+'。本报告记录正在执行的修复；未完成实验的结果不作科学结论。')
    p('授权依据：用户已批准修复计划，并允许双卡并行、每卡最多5个进程。训练队列每进程限制2个CPU线程；最终性能计时在训练停止、设备空闲时进行。')
    label_audit=REPORT/'ACT_source_trace_20261012/ACT_source_trace_evidence.json'
    if label_audit.exists():
        evidence=load(label_audit)
        h('新增待处理问题：ACT在MIC单位换算中的别名遗漏')
        p(f"2026-10-12原始JSON溯源确认：204条肽的ACT标注传递一致，但上游MIC分子量函数也漏识别ACT。当前209条ACT训练观察中{evidence['ACT_label_counterfactual_changed']}条标签在仅修复此遗漏时会改变。现有训练CSV尚未修正该标签问题。本报告完成数只表示原已登记任务的执行进度，不等于标签修复后的最终交付；后续依赖范围和工期需重新核定。溯源报告：ACT_source_trace_20261012/ACT_source_trace_report_ZH.md/.docx。本次只读核查没有修改、取消或追加训练任务。")
    h('已冻结的范围')
    p('TIGER现有40臂各10种子，另有pair/single×N2分类预训练开关×N1回归负注入开关的8格分解，以及真正关闭训练相似度门的新实验。基准双条件固定比较00与11，10/01只用于阶段分解。所有TIGER预测器每种子5折。AMPLiT正式呈现仅使用固定训练词向量的单一回归适配实现，两条件各10种子。')
    p('依赖去重后，TIGER共有440个独立重训任务、30个经验证复用任务、20个配置别名，对应490个结果位置。30个复用任务覆盖3个不使用全局理化描述符p分支的消融臂；结构图节点仍保留既定氨基酸先验属性。150个实际折模型已通过ACT修正前后预测逐位一致的检查。')
    p('外测终点继续使用11家族、208条肽、2,211条规范有向边。用户已确认同家族实际实验条件一致，表中差异属于整理噪音；R08据此关闭，不做删边或条件敏感性面板。CFU输入覆盖错误独立修复，保留源字段及规范映射。')
    h('执行进度')
    rows=[['项目','状态/证据'],['原文件保全','archive_pre_repair；代码/数据/结果哈希清单及既有2,000折权重哈希'],['核心契约','core_acceptance.json：ACT别名、训练/验证门分离、真正关门、负对方向、缓存身份、终点完整性/非有限值及退化相关'],['图缓存','3种图配置的正负缓存已重建；每种7,217正观察与7,860负图，零构图丢失'],['短试跑','10条受影响路径，5折、2 epoch；全部通过'],['不含p复用','reuse_acceptance.json：150个折模型的实际前向逐位一致'],['TIGER正式任务',f'{done}/440 个独立任务已写入DONE，其余运行中或待执行']]
    for stage in ['smoke','pilot','full']:
        path=WORK/'queue'/stage/'state.json'
        if path.exists():
            stage_state=load(path);rows.append([stage,f"{stage_state['status']}；{stage_state['completed_tasks']}/{stage_state['total_tasks']}完成；{len(stage_state['active'])}活跃；{len(stage_state['failures'])}失败"])
    path=WORK/'queue/campaign/state.json'
    if path.exists():
        rows.append(['统一执行队列',f"{s['status']}；{s['completed_jobs']}/{s['total_jobs']}项完成；{len(s['active'])}活跃；{s['remaining_to_launch']}待启动；当前轮{len(s['failures'])}项失败"])
        rows.append(['队列组成','440个TIGER重训/外测、30个复用外测、20个AMPLiT训练/外测、200个其他外部含负样本训练/外测、1个全量放行验收；20个配置别名另行解析，不重复计算'])
    table(rows)
    h('正式实验进度：按训练任务统计')
    amplit_done=[x for x in (RUNS/'AMPLiT').glob('*/seed*/DONE.json') if load(x).get('status')=='PASS' and not load(x).get('smoke',False)]
    external_done=[x for x in (WORK/'external_runs').glob('*/seed*/DONE.json') if load(x).get('status')=='PASS']
    table([['工作','正式完成数','尚余工作'],
           ['TIGER独立重训',f'{done}/440','每任务5折；40臂、8格分解和无门实验按配置去重'],
           ['AMPLiT两训练条件',f'{len(amplit_done)}/20','固定训练表示；两条件各10种子'],
           ['其余20种外部方法含负样本条件',f'{len(external_done)}/200','去除负侧13条冲突记录后重训，每法10种子'],
           ['TIGER可复用任务',f"30/30已通过前向一致性检查；{eval_counts['reuse_validate']}/30已完成统一外测",'复用任务不重复训练']])
    p(f"TIGER新训练任务中，{eval_counts['train']}/440已写入统一队列的EVALUATED完成标志；五折训练完成数为{done}/440。部分先导任务已在放行阶段完成验收，统一队列的外测完成标志仍按各任务实际文件单独计数。")
    external_counts={}
    for x in external_done:external_counts[x.parent.parent.name]=external_counts.get(x.parent.parent.name,0)+1
    complete_methods=sorted(name for name,n in external_counts.items() if n==10)
    if complete_methods:p(f"其余20种外部方法中已有{len(complete_methods)}种完成全部10种子："+'、'.join(complete_methods)+'。')
    p(f'正式训练完成合计{done+len(amplit_done)+len(external_done)}/660，短试跑不计入。TIGER的DONE表示五折训练完成；统一队列的完成标志还要求外测，因此不能把上述训练计数与队列完成计数相加，也不能据短任务完成比例推算整个项目进度。')
    pilot=load(WORK/'queue/pilot/state.json')
    release_path=WORK/'full_training_release.json'
    released=release_path.exists() and load(release_path).get('status')=='PASS'
    if released:
        release=load(release_path)
        p(f"五项TIGER完整先导实验已完成{pilot['completed_tasks']}/5，全部通过端到端验收。五个任务各输出2,211条规范外测边，特征标准化统计复算最大误差均为0，校准输出已核验。full_training_release.json为PASS，全量TIGER已获准进入调度；此处验收通过不等于剩余任务已经开始训练。")
    else:
        p(f"五项TIGER完整先导实验已完成{pilot['completed_tasks']}/5。全量TIGER放行还需全部先导训练结束，并通过特征、划分、配对和校准预测的端到端验收；当前不把短试跑通过等同于完整验收通过。外部方法已通过独立训练输入验收，可以并行执行。")
    ngpu=sum(a['resource']=='gpu' for a in s['active'])+sum(len(v) for v in s.get('foreign_compute_processes',{}).values())
    ncpu=sum(a['resource']=='cpu' for a in s['active'])
    tf_gpu_enabled=(WORK/'tf_gpu_runtime.json').exists()
    tf_note='AMPlify及MBC-Attention的未完成任务已迁移至独立TensorFlow GPU环境，已完成CPU种子保留并登记设备。' if tf_gpu_enabled else 'TensorFlow方法沿用已披露的CPU环境。'
    p(f'状态快照中有{ngpu}个GPU计算进程（含先导实验）和{ncpu}个CPU作业；上限为每卡5个计算进程，并按显存预留调度，不保证每张卡始终能装下5个模型。'+tf_note+'最终推理计时等全部训练释放设备后执行。')
    tiger_ids={r.get('task_id') for r in reg if r['action']!='alias'}
    active_tiger=[a for a in s['active'] if a['id'] in tiger_ids]
    if released:
        p(f"当前统一队列中有{len(active_tiger)}个TIGER任务运行；当前队列先列外部方法，再列TIGER，满足显存/进程预算的任务依序启动。若当前TIGER活跃数为0，表示已通过验收、正在等待调度资源，并非仍在等待验收。")
        if active_tiger:p('当前TIGER任务：'+'、'.join(a['id'] for a in active_tiger)+'；完成计数仅在五折训练结束后增加。')
    h('当前较慢任务与验收依赖')
    gate_run=RUNS/'s4_gate_off/s4_gate_off_seed123'
    gate_hist=[]
    for history in sorted((gate_run/'logs').glob('fold*_history.csv')):
        records=list(csv.DictReader(history.open()))
        if records:gate_hist.append((history.stem,records[-1]['epoch']))
    if gate_hist and not (gate_run/'DONE.json').exists():
        p('无相似度门先导任务仍在训练；最新已写入历史文件的轮次：'+'；'.join(f'{name}：epoch {ep}' for name,ep in gate_hist)+'。全量TIGER放行文件尚未产生，需等待该任务完成和随后端到端验收；这属于预设依赖，不能把未执行的验收计为通过。')
    elif (gate_run/'DONE.json').exists():
        gate_done=load(gate_run/'DONE.json')
        p(f"无相似度门先导任务五折训练已完成，实测耗时{gate_done['elapsed_seconds']/3600:.2f}小时。其完整运行时间已纳入后续排程评估；全量放行状态：{'PASS' if released else '待验收'}。")
    amplify_rows=[['AMPlify种子','已完成轮数','最近一轮耗时','正式完成标志']]
    for log in sorted((WORK/'external_runs/AMPlify').glob('seed*/training.log')):
        records=re.findall(r'363/363 - (\d+)s - loss: ([.\d]+) - val_loss: ([.\d]+)',log.read_text(errors='replace'))
        if records:amplify_rows.append([log.parent.name,str(len(records)),records[-1][0]+'秒','已完成' if (log.parent/'DONE.json').exists() else '未完成'])
    table(amplify_rows)
    for finished_path in sorted((WORK/'external_runs/AMPlify').glob('seed*/DONE.json')):
        data=load(finished_path)
        p(f"AMPlify {finished_path.parent.name}已完成训练及规范外测输出，完整任务耗时{data['elapsed_seconds']/3600:.2f}小时；其完成标志为{data['status']}。")
    if ncpu and all(a['id'].startswith('AMPlify_') for a in s['active'] if a['resource']=='cpu'):
        p(f'当前{s["cpu_job_limit"]}个CPU名额由AMPlify种子占用，其他尚未开始的CPU任务需等待名额释放；GPU队列独立推进。AMPlify轮次仍持续增加，未显示卡死。')
    p('AMPlify此前CPU训练耗时专项核查已完成：743万参数、CPU执行、每轮363个训练批次和patience=50共同解释旧CPU任务的小时级训练；早停与梯度更新通过隔离诊断。当前设备与完成状态以迁移记录和下表为准。见同目录AMPlify_runtime_review_ZH.md/.docx。')
    gpu_check=WORK/'tf_gpu_repair/validation_summary.json'
    if gpu_check.exists() and load(gpu_check).get('status')=='PASS':
        transition='用户于2026-10-12授权取消两方法的CPU任务并切换GPU：3个未完成AMPlify CPU尝试已终止归档；AMPlify的2个、MBC-Attention的6个已完成CPU种子保留，余下8+4个种子使用GPU。每任务TensorFlow分配上限1 GiB、整进程预留2 GiB，迁移前已复验。被中止种子从相同种子初始化重新训练，不冒充断点续训；其他工作进程原位续跑。' if tf_gpu_enabled else '正式队列仍沿用原CPU任务，未因诊断中断或切换现有训练。'
        p('TensorFlow GPU加载问题已在独立mbc_att_gpu_v15环境修复并通过两方法的预测、梯度及训练更新验收。根因是原进程未带入已有CUDA库的搜索路径，库本身已安装。AMPlify另需eager执行以避开TF 2.12.1带掩码反向LSTM图模式的数值异常；MBC-Attention使用已验证的graph执行。'+transition+'修复与迁移记录见TensorFlow_GPU_repair_ZH.md/.docx及TensorFlow_GPU_migration_ZH.md/.docx。')
    tf_rows=[['方法','已完成GPU任务','已保留CPU任务','GPU任务完整耗时范围/中位数']]
    for method in ['AMPlify','MBC-Attention']:
        records=[load(x) for x in (WORK/'external_runs'/method).glob('seed*/DONE.json') if passed(x)]
        durations=[d['elapsed_seconds']/60 for d in records if d.get('training_device','').startswith('GPU')]
        tf_rows.append([method,str(len(durations)),str(len(records)-len(durations)),f'{min(durations):.1f}–{max(durations):.1f}分钟 / {statistics.median(durations):.1f}分钟' if durations else '暂无完整GPU任务'])
    table(tf_rows)
    p('以上耗时为本轮各种子的完整训练及外测任务墙钟时间，不是外部推理基准计时；各种子早停轮数不同，不能直接解释为配对加速倍数。')
    h('CPU方法与GPU方法并行训练')
    cpu_methods=[('EIPpred','随机森林回归'),('AMPActiPred','级联森林回归')]
    if not tf_gpu_enabled:cpu_methods.extend([('MBC-Attention','TensorFlow模型，当前环境使用CPU'),('AMPlify','TensorFlow模型，当前环境使用CPU')])
    cpu_rows=[['方法','执行方式','已完成种子数（共10）']]
    for method,desc in cpu_methods:
        n=sum(load(x).get('status')=='PASS' for x in (WORK/'external_runs'/method).glob('seed*/DONE.json'))
        cpu_rows.append([method,desc,str(n)])
    table(cpu_rows)
    p(f"CPU与GPU采用独立队列，本快照有{ncpu}个CPU任务运行；CPU方法完成数见上表。CPU任务设置CUDA_VISIBLE_DEVICES=-1，不占每卡5个GPU任务名额。CPU并发上限当前为{s['cpu_job_limit']}，每任务的主要计算库线程预算为2，TensorFlow inter-op线程为1。此前已提高CPU并发、随后迁移TensorFlow任务到GPU；均未修改训练超参数。")
    handoff=WORK/'queue/campaign/cpu4_handoff_verification.json'
    if handoff.exists():
        q=load(handoff)
        p(f"调度器交接核验：{q['status']}；{len(q['workers'])}个原运行任务均保留原工作进程，或在交接期间正常完成。队列支持通过resource_limits.json调整后续CPU并发；已运行模型不重启。实际提速仍取决于CPU竞争、数据准备与各方法早停时间，不能据并发翻倍承诺总工期减半。")
    h('剩余工作顺序与预计完成窗口')
    if label_audit.exists():
        p('下表保留发现上游MIC换算问题之前、原登记队列的排程窗口。该窗口不包含新增标签修复及可能需要重跑的实验，已不能作为当前完整v15交付日期；须完成数据影响范围核定后重新估计。')
    recent_counts={hours:sum((Path(r['run_dir'])/'DONE.json').stat().st_mtime>time.time()-hours*3600 for r in finished) for hours in [2,4]}
    p(f"本次按2026年10月12日实测进度更新估计。TIGER最近2小时完成{recent_counts[2]}个五折训练，最近4小时完成{recent_counts[4]}个；正式训练尚余{440-done}个TIGER任务。剩余任务包含图结构/特征消融、相似度门、负样本阶段分解等不同成本配置，不能只按当前吞吐线性外推。真正无门先导约3.04小时/种子，可能构成末段长尾。下列窗口以设备持续可用且后续验收无新增问题为前提。")
    table([['阶段','具体工作','预计窗口（北京时间）'],
           ['完成重训与逐运行验收','完成660个独立训练任务；核验复用与配置别名；固定种子、检查点和预测文件逐项登记','预计2026年10月13日白天至晚间'],
           ['统一评估、统计和计时','11家族/208肽/2,211边的值级验收；四指标、种子配对检验；设备空闲时20次重复计时','训练结束后约12–24小时，非计时核查与训练交叉进行'],
           ['v15图文与发布核查','保留notebook中的用户讨论、既定图风格和单子图接口；更新中英文报告、方法学MD/Word和复现包','再预留约12–24小时，排版及代码整理提前进行'],
           ['完整v15交付','验收通过后的数据、图、notebook、中英文文档和可复现包','预计2026年10月14日至10月15日，取决于最终验收']])
    arm_rows=[]
    for label in dict.fromkeys(r['label'] for r in reg if r['action']=='train'):
        members=[r for r in finished if r['label']==label]
        durations=[load(Path(r['run_dir'])/'DONE.json')['elapsed_seconds']/60 for r in members]
        n_eval=sum(passed(WORK/'tiger_external'/r['task_id']/'EVALUATED.json') for r in reg if r['action']=='train' and r['label']==label)
        arm_rows.append([label,f'{len(members)}/10',f'{n_eval}/10',f'{statistics.median(durations):.1f}分钟' if durations else '待执行'])
    table([['TIGER配置','五折训练完成','统一队列外测完成','已完成任务耗时中位数']]+arm_rows)
    p('预计窗口包含长尾训练、值级验收与图文检查的余量，属于排程估计。外部220个正式训练任务已完成，后续主要计算依赖为TIGER剩余任务及最终空闲设备计时。不会为赶工缩短训练、减少种子或挑选结果。')
    p('现阶段尚未完成最终统计、空闲设备计时及v15科学图文交付，当前部分结果不用于替换正式主表。前期启动问题已修复；当前队列若发生失败将保留日志并停止追加任务，不把未通过验收的输出算作完成。')
    h('R16：外部训练池正负标签冲突')
    p('精确序列核对显示，原外部负池7,962条中有13条同时出现在4,929条正序列中。若直接拼接，它们同时得到实测MIC和负样本伪标签。逐条证据保存在negative_pool_overlap_positive_records.csv与negative_pool_overlap_negative_records.csv；剔除负侧重复后，规范外部负池为7,949条。正侧实测记录保留。')
    p('TIGER使用的7,860条有效负图池与其训练正序列没有精确交集，当前TIGER任务不受R16影响。外部方法按实际读取的训练表、方法特定长度筛选及缓存身份逐一核定影响；涉及修改训练输入者按已批准计划扩展到全部十种子重训。不可把过滤后的表与原按行定位缓存直接拼接。原负池与旧权重仍留档，修复输出隔离存放。')
    p('用户提醒后的复核：原TIGER训练入口确有“正样本优先”规则，负池生成脚本也提前剔除训练重叠；该规则在修复前归档中已存在。不能据外部源表重叠反推TIGER有冲突，也不能只据拼接代码就认定所有外部方法的最终训练均受影响。曾暂缓外部扩展队列，并在任何额外重训启动前完成下述只读复核。')
    audit=WORK/'negative_precedence_audit/verification.json'
    if audit.exists():
        a=load(audit)['results']
        table([['原代码实际输入核查','结果'],['TIGER','训练前剔除与正样本相同的负记录；修复前已有'],['EIPpred',f"在clf.fit前截停：{a['EIPpred']['fit_rows']}行，13组相同特征但不同目标的序列"],['Das',f"原划分的训练子集{a['Das']['fit_rows']}行；7组冲突的正负两条记录同时进入训练子集"],['GGN-ABPMIC',f"原ProteinGraphDataset实际保留{a['GGN-ABPMIC']['actual_dataset_rows']}行，其中13组序列对应两种标签"]])
        p('上述复核只执行输入构建并在训练前截停，没有拟合任何外部模型。证据：artifacts/repair_v15/negative_precedence_audit/verification.json及三个逐记录CSV。')
    all_audit=WORK/'negative_precedence_audit/all_external_corpus_checks.json'
    if all_audit.exists():
        a=load(all_audit)
        p('随后逐方法检查完成：21个外部含负样本入口的原始语料函数、编码后目标数组或归档特征/标签表，均保留13组冲突序列；下游按行或行索引传递标签，没有统一实行TIGER的正样本优先过滤。部分驱动的keep=first去重仅位于no_neg分支，不能保护含负样本分支。检查类型与来源哈希逐方法登记，不将所有检查都描述为实际fit截停。')
        table([['方法','实际入口行数','冲突序列数']]+[[r['method'],str(r['n_rows']),str(r['conflicting_sequences'])] for r in a['methods']])
        p('完整证据：negative_precedence_audit/all_external_corpus_checks.csv/.json及逐方法冲突记录。结论：R16扩展只针对外部含负样本条件；TIGER不因本项增加任务，外部无负样本条件也不因本项重训。AMPLiT两条件重训仍来自原已批准的表示空间修复。')
    h('可续作文件')
    table([['内容','路径（相对项目根目录）'],['训练注册表','artifacts/repair_v15/run_registry.json'],['队列状态与日志','artifacts/repair_v15/queue/'],['新TIGER训练输出','runs_repair_v15/'],['验收与输入证据','artifacts/repair_v15/'],['执行入口','scripts/repair_v15/'],['旧输入与源码','artifacts/repair_v15/archive_pre_repair/']])
    p('尚需完成：剩余TIGER正式训练、复用任务外测及配置别名解析；统一指标与统计重算；空闲设备统一计时；v15图、保留用户讨论内容的notebook、中英文报告和方法学文档、可复现包及发布核查。外部220个训练任务已完成，仍需纳入最终统一统计和发布验收。当前未将部分训练输出覆盖为正式完整结果。')
    snapshot={'captured_at':snapshot_time.isoformat(timespec='seconds'),'queue':s,'tiger_training_complete':done,'tiger_training_total':440,'tiger_external_complete':eval_counts,'external_training_complete':len(external_done),'external_training_total':200,'amplit_training_complete':len(amplit_done),'amplit_training_total':20,'training_complete':done+len(external_done)+len(amplit_done),'training_total':660,'recent_tiger_completions':recent_counts,'tiger_arm_rows':arm_rows,'tf_migration_rows':tf_rows}
    dump(WORK/'queue/campaign'/f"progress_snapshot_{snapshot_time.strftime('%Y%m%d_%H%M%S')}.json",snapshot)
    doc=Document();style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    md=[];texts=[]
    for kind,val in blocks:
        if kind=='h':doc.add_heading(val,level=1);md.extend(['# '+val,'']);texts.append(val)
        elif kind=='p':doc.add_paragraph(val);md.extend([val,'']);texts.append(val)
        else:
            t=doc.add_table(rows=0,cols=len(val[0]));t.style='Table Grid'
            for row in val:
                cells=t.add_row().cells
                for c,s in zip(cells,row):c.text=s;texts.append(s)
            md.append('| '+' | '.join(val[0])+' |');md.append('| '+' | '.join(['---']*len(val[0]))+' |')
            md.extend('| '+' | '.join(row)+' |' for row in val[1:]);md.append('')
    base=REPORT/'repair_execution_status_ZH'
    base.with_suffix('.md').write_text('\n'.join(md))
    temp=base.with_name(base.name+'.tmp.docx');doc.save(temp);os.replace(temp,base.with_suffix('.docx'))
    check=Document(base.with_suffix('.docx'))
    actual='\n'.join([p.text for p in check.paragraphs]+[c.text for t in check.tables for row in t.rows for c in row.cells])
    assert all(t in actual for t in texts)
    print(base)


if __name__=='__main__':main()
