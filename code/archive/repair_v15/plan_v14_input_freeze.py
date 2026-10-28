"""Read-only task reconciliation after the user froze V14 actual inputs."""
from datetime import datetime
from zoneinfo import ZoneInfo
from collections import Counter
from common import ROOT, WORK, REPORT, load, dump, sha, run_config
import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from code.utils.config import load_config
from registry import effective


def main():
    out=REPORT/'V14_input_freeze_20261012';out.mkdir(parents=True,exist_ok=True)
    reg=load(WORK/'run_registry.json')
    state=load(WORK/'queue/campaign/state.json')
    moment=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(timespec='seconds')
    gates={'S4_sim00','S4_sim03','S4_sim05','S4_sim09'}
    newfactors={'factor_P00','factor_P01','factor_P11','factor_S00','factor_S01'}
    rows=[];references=[]
    for r in reg:
        item=dict(label=r['label'],seed=r['seed'],group=r['group'],old_action=r['action'])
        if r['group']=='ablation40':
            action='train_with_V14_inputs' if r['label'] in gates else 'reuse_V14_weights_after_acceptance'
            reason='Fix validation gate at 0.7; training gate intervention remains' if r['label'] in gates else 'ACT feature repair withdrawn; historical training protocol retained'
            original=run_config(r['label'],r['seed']).parent
        elif r['group']=='gate_off':
            action='train_with_V14_inputs';reason='True disabled training gate, validation gate fixed at 0.7';original=None
        else:
            action='train_with_V14_inputs' if r['label'] in newfactors else 'reuse_V14_weights_after_acceptance'
            reason='Missing factorial condition or joint-pair direction bug' if r['label'] in newfactors else 'Existing V14 factorial-equivalent predictor'
            arm={'factor_P10':'S4','factor_S11':'s4_single_n1n2','factor_S10':'s4_single_n2_noneg'}.get(r['label'])
            original=ROOT/f'runs_ablation/{arm}/{arm}_seed{r["seed"]}' if arm else None
        item.update(proposed_action=action,reason=reason,old_run=r.get('run_dir',''),reuse_run=str(original) if action.startswith('reuse') else '')
        if action.startswith('reuse'):
            assert original is not None
            required=[original/'config.json',original/'results/calibrator.json']+[original/'checkpoints'/f'fold{f}_best.pt' for f in range(1,6)]
            assert all(p.exists() for p in required),original
            references.append({'label':r['label'],'seed':r['seed'],'run':str(original),'artifacts_present':True,'config_sha256':sha(original/'config.json')})
            if r['label']=='factor_S10':
                old=load_config(original/'config.json').to_dict();old.update(validation_similarity_threshold=old['similarity_threshold'],train_similarity_gate_enabled=True)
                assert effective(old)==effective(load(r['config_path']))
        if 'run_dir' in r:
            done=ROOT/r['run_dir']/'DONE.json'
            item['ACT_version_training_complete']=done.exists() and load(done).get('status')=='PASS'
        rows.append(item)
    table=pd.DataFrame(rows);table.to_csv(out/'proposed_task_disposition.csv',index=False,encoding='utf-8-sig')
    assert len(table)==490
    assert sum(table.proposed_action.eq('train_with_V14_inputs'))==100
    old_cancel=table[(table.old_action=='train') & table.proposed_action.str.startswith('reuse')]
    assert len(old_cancel)==340
    source=ROOT/'metadata/train_val_by_cfu_group_ug_per_mL.csv'
    snapshot=ROOT/'reports/S4_study_nature_v14/source_snapshot/methods_package_S4_v1/1_training_data_pairing/data/train_val_by_cfu_group_ug_per_mL.csv'
    pos=ROOT/'reports/methods_package_S4_v1/4_S4_benchmark/data/train_restricted4929_tiger.csv'
    oldpos=WORK/'archive_pre_repair/reports/methods_package_S4_v1/4_S4_benchmark/data/train_restricted4929_tiger.csv'
    assert sha(source)==sha(snapshot) and sha(pos)==sha(oldpos)
    external_done=[p for p in (WORK/'external_runs').glob('*/seed*/DONE.json') if load(p).get('status')=='PASS']
    amplit_done=[p for p in (ROOT/'runs_repair_v15/AMPLiT').glob('*/seed*/DONE.json') if load(p).get('status')=='PASS']
    assert len(external_done)==200 and len(amplit_done)==20
    done=sum(bool(r.get('ACT_version_training_complete')) for r in rows if r['old_action']=='train')
    evidence={'status':'PROPOSED_SCOPE_RECONCILED_NOT_EXECUTED','time':moment,
              'user_scope':'V14 labels, terminal annotations, terminal feature calculation and external CFU encoding frozen',
              'new_TIGER_training_tasks':100,'TIGER_folds':500,'ablation_reuse_result_slots':360,
              'factorial_reuse_slots':30,'old_training_requirements_eliminated':340,
              'external_completed_retained':200,'AMPLiT_completed_retained':20,
              'total_retained_repair_training_tasks':320,'old_TIGER_ACT_training_complete':done,
              'old_queue_snapshot':state,'reuse_reference_evidence':references,
              'raw_positive_input_hashes':{str(p):sha(p) for p in [source,snapshot,pos,oldpos]},
              'limits':'Reuse requires final checkpoint/prediction provenance acceptance; no GPU re-inference in this scope audit. Negative-conflict filtering remains an explicit non-chemical training-corpus correction.',
              'queue_or_training_code_modified':False}
    dump(out/'scope_evidence.json',evidence)

    blocks=[]
    def h(t):blocks.append(('h',t))
    def p(t):blocks.append(('p',t))
    def t(rows):blocks.append(('t',[[str(x) for x in row] for row in rows]))
    h('冻结V14实际输入后的任务缩减与保留计划')
    p('核定时间：'+moment+'。用户明确确认：以V14训练、测试数据为准，并保持V14实际模型输入，包括当时的末端特征计算和外测CFU编码。本文件给出任务调整清单；没有向调度器发送停止信号，也没有切换正在运行的代码。')
    p('结论：撤下ACT特征、ACT分子量/MIC重算、AM→AMD再整理和外测CFU改码驱动的任务扩张。按保留已批准的独立方法修复计算，TIGER仍需100个种子任务、每任务5折，共500折；外部已完成的200个含负修复任务及AMPLiT两条件20个任务保留。修复训练总规模由原660降为320，其中220已完成、100需在V14实际输入下登记执行；此前890任务的标签修复方案不再适用。')
    h('一、数据冻结边界')
    p('保持V14序列、正样本MIC标签、末端字段、末端特征计算、折内标准化配方，以及外测CFU编码。外测仍为11家族、208条肽、2,211条规范有向边。当前TIGER源CSV与V14快照逐字节相同，外部正样本表与修复前归档逐字节相同；本轮迄今没有应用ACT质量修正到生产标签。改变的是TIGER特征代码及部分外测输入，因此只保留CSV并不足以恢复V14输入。')
    p('保留一项明确的、与末端化学无关的语料修复：外部含负训练仍实行正样本优先，从V14来源负池剔除13条与正序列冲突的负记录，7,962→7,949；正样本序列与MIC保持V14。这意味着外部含负最终fit语料并非逐行原封不动的V14拼接表。该项属于此前已核实并批准的冲突修复，不能被写成化学信息补全，也不能在方法学中隐去。TIGER原有效负池已无该精确冲突，无需因本项重训。')
    p('末端空值按V14既有实现处理，并在Methods中披露其信息不完整及默认编码假设；不能声称空值已证实为游离端，也不能将冻结数据表述为数据库化学标注全部正确。ACT与AM/AMD发现留在审计记录，本轮不据此更改输入。外测CFU表述改为V14实际统一编码，而不再声称模型使用全部逐肽真实CFU。')
    h('二、40臂消融：36臂复用、4臂重训')
    t([['类别','结果位置/训练数','处理'],['原40臂中的36臂','360个种子结果','复用V14检查点、匹配标准化统计与校准器，统一复核外测'],['相似度0、0.3、0.5、0.9四臂','40个新种子任务','用V14输入；训练阈值各异，验证阈值统一固定0.7'],['真正无训练相似度门','10个新种子任务','跳过训练相似度判断，验证阈值仍固定0.7']])
    p('原ACT特征修复要求重训37臂、370个种子；其中33臂、330个种子不再需要重训，4个阈值臂仍因验证规则独立修复而保留。另3个无p臂原本就只做复用核验。完整TIGER-pair主模型包含在可复用36臂内，无须因为化学标注不完整重训主模型。')
    p('四阈值臂不能直接作为新的单因素门控比较沿用旧权重：V14同时改变训练配对阈值和验证邻域，影响选模与校准。当前保留已批准的“固定验证规则、只改变训练门”的科学设计，因此仍须重做这四组。旧0.0结果只能称阈值0.0；真正无门独立列组。基线0.7的V14模型可复用。')
    p('S4_negpairs为旧top-k负对臂（topk=1），不等于joint负对的P11。代码对照及小型真实函数探针显示，保留top-k配方时其返回有向对列表不受本轮joint方向修复影响，可纳入36臂复用；P11的joint方向修复仍需训练。')
    h('三、八格负样本分解：三格复用、五格新训')
    t([['配置','分类预训练用负样本','回归用负样本','处理'],['P00','否','否','新训10种子'],['P10','是','否','复用V14完整TIGER-pair'],['P01','否','是','新训10种子；使用修正的joint方向规则'],['P11','是','是','重训10种子；旧joint结果受方向判门问题影响'],['S00','否','否','新训10种子'],['S10','是','否','复用V14 s4_single_n2_noneg，已有10种子'],['S01','否','是','新训10种子'],['S11','是','是','复用V14 s4_single_n1n2']])
    p('P=pair，S=single。八格仍保留80个结果位置，只有50个需要新训练；P10与S11已含在原40臂，S10位于40臂之外。此次核对S10的有效训练配置与计划一致，并找到全部10种子的5折检查点及校准器；无效的single负对模式字段差异不构成新训练。旧表将S10称“无负”只能指回归阶段，它不能替代全流程不使用负样本的S00。正式复用仍需通过检查点来源与数值复现验收。')
    p('总计：40个阈值任务＋10个真正无门任务＋50个因子任务＝100个TIGER独立种子任务。八格00/11的基准主比较沿用原登记；10/01保留作阶段分解，不按外测分数重新挑选分组。')
    h('四、外部方法：保留220个已完成修复任务')
    t([['工作','任务数','当前处理'],['AMPLiT含负/无负','2×10＝20，已完成','保留；训练词向量固定、查询不重新拟合，两条件一致'],['其余20种外部方法含负','20×10＝200，已完成','保留；V14正标签＋正样本优先的负池过滤'],['其余20种外部方法无负','原有200个结果','复用V14；撤下因ACT标签修正提出的再次训练']])
    p('外部220个完成标志均为PASS，其训练读取的正样本表仍是V14标签。因此无需因撤回ACT分子量修正再训练它们。AMPlify和MBC-Attention已经完成的GPU/CPU种子及环境修复保留。以上训练数包含AMPLiT含负10种子，不与其他外部含负任务重复计数。')
    h('五、当前ACT版队列与下一步')
    p(f'本次快照：原ACT版TIGER完成{done}/440个五折训练；统一队列完成{state["completed_jobs"]}/{state["total_jobs"]}项，活跃{len(state["active"])}项、待启动{state["remaining_to_launch"]}项。队列项包括外测与验收，不能与独立训练数相加。')
    p('范围上消除340个原登记TIGER重训需求：330个既有消融种子＋10个可复用S10种子。保留的100个任务也必须使用V14实际输入重新登记，不能让当前ACT特征版的同名任务直接续跑后混入V14比较。已完成ACT版TIGER结果保留为审计分支；旧V14检查点继续作为正式复用来源。原先因ACT质量问题提出的额外30个无p任务和200个外部无负任务不再追加。')
    p('执行顺序：先停发并收束旧ACT版TIGER队列；在独立分支封装V14末端特征与外测CFU输入，同时保留指标、缓存身份、无门及joint方向等必要修复；以V14基线做数值复现验收；登记100个必要训练，按双卡每卡最多5进程并遵守显存预算执行；统一评估、统计、空闲设备计时和图文更新。不要在旧进程仍运行时原地改公共特征函数。')
    p('仍需完成的非训练工作：固定11家族/2,211边的完整性和退化相关处理；消融目录/种子完整性核验；末端与CFU实际输入的Methods更正；相似度定义重算与披露；Pareto事后最优选择的统计边界说明；可复现依赖、统一计时和最终图文。冻结V14输入不意味着恢复已证实有逻辑缺陷的评估器。')
    p('交付：proposed_task_disposition.csv列出490个TIGER结果位置的建议处理；scope_evidence.json保存当前队列快照、输入哈希和复用配置证据。复现代码为scripts/repair_v15/plan_v14_input_freeze.py。本报告是具体调整方案，尚未执行取消/调度切换。')

    doc=Document();style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10);style.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    md=[];texts=[]
    for kind,val in blocks:
        if kind=='h':doc.add_heading(val,level=1);md.extend(['# '+val,'']);texts.append(val)
        elif kind=='p':doc.add_paragraph(val);md.extend([val,'']);texts.append(val)
        else:
            tab=doc.add_table(rows=0,cols=len(val[0]));tab.style='Table Grid'
            for row in val:
                for cell,text in zip(tab.add_row().cells,row):cell.text=text;texts.append(text)
            md.extend(['| '+' | '.join(val[0])+' |','| '+' | '.join(['---']*len(val[0]))+' |'])
            md.extend('| '+' | '.join(row)+' |' for row in val[1:]);md.append('')
    base=out/'V14_input_freeze_plan_ZH';base.with_suffix('.md').write_text('\n'.join(md));doc.save(base.with_suffix('.docx'))
    check=Document(base.with_suffix('.docx'));content='\n'.join([p.text for p in check.paragraphs]+[c.text for t in check.tables for row in t.rows for c in row.cells])
    assert all(text in content for text in texts)
    print({'status':evidence['status'],'time':moment,'TIGER_new':100,'external_retained':220,'old_ACT_TIGER_done':done,'active':state['active'],'output':str(base)})


if __name__=='__main__':main()
