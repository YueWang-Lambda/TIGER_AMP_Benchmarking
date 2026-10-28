"""Matching Markdown/Word live status for the approved V14-input transition."""
from datetime import datetime
from zoneinfo import ZoneInfo
import shutil
from common import *
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn


def main():
    reg=load(WORK/'run_registry.json');state=load(WORK/'queue/state.json')
    cancel=load(ROOT/'artifacts/repair_v15/v14_input_freeze/cancellation.json')
    acceptance=load(WORK/'input_and_replay_acceptance.json')
    def passed(p):return p.exists() and load(p).get('status')=='PASS'
    trained=[r for r in reg if r['action']=='train' and passed(Path(r['run_dir'])/'DONE.json')]
    inferred={k:sum(passed(WORK/'tiger_external'/r['task_id']/'EVALUATED.json') for r in reg if r['action']==k) for k in ['train','reuse']}
    smoke=sum(passed(WORK/'smoke'/f'{label}_seed123/DONE.json') for label in ['factor_P00','factor_P11','factor_S00','S4_sim03','s4_gate_off'])
    release=passed(WORK/'full_training_release.json')
    retained=load(WORK/'retained_external_jobs.json')
    assert len(retained)==220
    assert all(sha(r['completion'])==r['sha256'] and passed(Path(r['completion'])) for r in retained)
    time=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(timespec='seconds')
    blocks=[]
    def h(t):blocks.append(('h',t))
    def p(t):blocks.append(('p',t))
    def table(rows):blocks.append(('t',[[str(c) for c in row] for row in rows]))
    h('v15执行进度：采用V14冻结实际输入的新分支')
    p('更新时间：'+time+'。用户已明确批准保留V14标签、末端特征计算和外测CFU编码，并要求立即取消不再需要的训练、启动必要任务。本文件取代原ACT特征修复队列的实时进度报告。')
    h('旧队列已取消，完成结果保留')
    p(f'在{cancel["time"]}冻结原调度器后，按进程组身份终止{len(cancel["terminated_process_groups"])}个正在运行的旧ACT版TIGER作业，取消后续派发。未强制KILL任何作业；收到SIGTERM的记录属于用户批准取消，不作为算法训练失败。原日志、未完成检查点和已完成ACT版结果保留为审计分支；旧调度及训练入口已添加禁止误重启保护。')
    p('其余20种外部方法含负条件200个任务及AMPLiT两条件20个任务完整保留，220个完成标志的哈希均未变化。外部含负训练继续保留此前已批准的正样本优先过滤：V14来源负池7,962→7,949，剔除13条与正序列冲突的负记录；V14正序列及MIC不变。此项是独立于化学标注的明确语料修复。')
    h('输入及数值复现验收')
    table([['核验','结果'],['运行时隔离','artifacts/repair_v15_v14input/runtime/code；不修改共享code中的特征实现'],['训练标签与末端字段','从V14来源冻结独立副本，ACT质量修正未写入生产标签'],['V14末端特征函数','与修复前归档features.py逐字节一致'],['正样本图输入','7,217条，模型消费张量与旧缓存全部相等'],['负样本图输入','7,860条，模型消费张量与旧缓存全部相等'],['五折理化标准化','复算与V14归档一致'],['外测输入','208条肽、11家族、2,211条边；CFU沿用V14统一编码1E5–1E6']])
    table([['V14 seed123重放','校准后逐边最大绝对差','宏PCC']]+[[r['method'],f"{r['max_abs_calibrated_error']:.9g}",f"{r['PCC']:.9f}"] for r in acceptance['replay']])
    p('以上数值为恢复旧检查点的输入/输出验收，不能当作本轮新训练成绩。相似度真关门、训练/验证门独立、joint负对方向、固定终点及非有限值/常数相关处理均通过独立探针。未知末端继续采用V14默认编码，不把空值解释为已证实游离端；这些数据局限将在方法学中披露。')
    if (WORK/'reuse_provenance_acceptance.json').exists():
        provenance=load(WORK/'reuse_provenance_acceptance.json');assert provenance['status']=='PASS'
        p(f"复用来源进一步核验：{provenance['runs']}个预测器的{provenance['checkpoints']}个折检查点及配置、标准化统计、校准器，共{provenance['checked_files']}个文件，与修复前/原V14归档哈希全部一致。该核查验证文件来源与完整性；所有复用模型仍按新严格评估入口逐一复评。")
    h('新任务范围及即时状态')
    table([['工作','总数','当前完成'],['TIGER相似度阈值0/0.3/0.5/0.9','40个种子任务','计入下方正式训练'],['真正无训练相似度门','10个种子任务','计入下方正式训练'],['负样本分解P00/P01/P11/S00/S01','50个种子任务','计入下方正式训练'],['正式TIGER五折训练','100',len(trained)],['正式新训练外测','100',inferred['train']],['V14权重复用外测','370个独立预测器',inferred['reuse']],['配置别名','20个结果位置','按注册表映射，不重复训练'],['AMPLiT及其他外部修复训练','220','220保留'],['短试跑','5条路径，每条5折/2 epoch',smoke]])
    p(f'当前队列阶段为{state["stage"]}，状态{state["status"]}，活跃作业{len(state["active"])}个；正式放行验收：'+('PASS。' if release else '等待短试跑全部完成后自动检查。'))
    if state['active']:table([['任务','GPU','PID']]+[[r['id'],r['gpu'],r['pid']] for r in state['active']])
    p('自动执行顺序：短试跑覆盖纯无负pair、joint负注入pair、纯无负single、阈值干预、真正无门五种路径；逐折输入/标准化/配置检查通过后自动启动100个正式任务，并在训练完成后导出规范外测预测；随后完成370个V14预测器的统一复评。保持双卡、每卡最多5个计算进程，每任务预留3 GiB并计入其他GPU作业，每进程2个CPU线程。任何失败停止新增派发，保留日志；恢复采用固定任务和种子，不按外测分数选择重试。')
    p('旧40臂中的36臂复用；四阈值臂需要固定验证门后重做。八格中的P10/S11与40臂复用重叠，S10另有10个V14预测器可复用，所以独立复用共370，别名结果位置20。所有正式TIGER训练仍为种子120–129、每种子5折，100任务对应500个折模型。')
    p('新范围为320个必要修复训练，其中220个外部任务已完成；100个TIGER任务在冻结V14输入的新分支执行。原660和后续890任务方案已被用户批准的新范围取代，不能再作为当前分母或工期依据。短试跑不计入正式完成数。')
    h('可续作入口与证据')
    table([['内容','路径（相对项目根目录）'],['新调度器','scripts/repair_v15_v14input/campaign.py --stage auto'],['新队列状态','artifacts/repair_v15_v14input/queue/state.json'],['新正式检查点','runs_repair_v15_v14input/'],['新外测预测','artifacts/repair_v15_v14input/tiger_external/'],['输入及旧值复现验收','artifacts/repair_v15_v14input/input_and_replay_acceptance.json'],['运行时/输入哈希','artifacts/repair_v15_v14input/freeze_manifest.json'],['原队列取消证据','artifacts/repair_v15/v14_input_freeze/cancellation.json'],['新训练/复用/别名登记','artifacts/repair_v15_v14input/run_registry.json'],['更新本报告','scripts/repair_v15_v14input/status_report.py']])
    p('首次调度入口命名为queue.py，与Python标准库重名，在任何训练启动前即报导入错误；已更名为campaign.py并重新冻结，首轮日志与manifest留存在revisions/launch_r1。模型源码、数据和验收结果未因此改变。')
    delivery=ROOT/'reports/S4_study_nature_v15/delivery_validation.json'
    if passed(delivery):
        h('V15分析与可视化已完成')
        p('正式交付已通过独立统计复算、23方法真实检查点统一计时、notebook执行、图文数值一致性和Word渲染检查。主目录为reports/S4_study_nature_v15；中英文报告为S4_study_report_zh/en.docx及同名Markdown，绘图入口为S4_complete_analysis_and_figures_Fig5_style.ipynb。原V14交付未覆盖，根目录中英基准/数据方法学已同步并保留更新前副本。')
        p('保留并披露的结果限制：AMPLiT含负seed122有一家族相关性未定义；V14实际末端特征和外测CFU编码按批准冻结；历史P10消融参考与00/11基准定义分开。详见V15_completion_audit_ZH。')
    else:
        p('训练后的剩余工作包括完整性/配置别名核对、四指标与配对统计、设备空闲时统一计时，以及按新口径更新图、notebook、中英文报告和方法学文件；训练/推理队列完成标记不代表全部论文交付完成。')
    doc=Document();style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10);style.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    md=[];texts=[]
    for kind,val in blocks:
        if kind=='h':doc.add_heading(val,level=1);md.extend(['# '+val,'']);texts.append(val)
        elif kind=='p':doc.add_paragraph(val);md.extend([val,'']);texts.append(val)
        else:
            tab=doc.add_table(rows=0,cols=len(val[0]));tab.style='Table Grid'
            for row in val:
                for cell,text in zip(tab.add_row().cells,row):cell.text=text;texts.append(text)
            md.extend(['| '+' | '.join(val[0])+' |','| '+' | '.join(['---']*len(val[0]))+' |']);md.extend('| '+' | '.join(row)+' |' for row in val[1:]);md.append('')
    REPORT.mkdir(parents=True,exist_ok=True);base=REPORT/'V14_input_freeze_execution_ZH'
    base.with_suffix('.md').write_text('\n'.join(md));doc.save(base.with_suffix('.docx'))
    check=Document(base.with_suffix('.docx'));content='\n'.join([p.text for p in check.paragraphs]+[c.text for t in check.tables for row in t.rows for c in row.cells]);assert all(t in content for t in texts)
    old=ROOT/'reports/repair_campaign_v15/repair_execution_status_ZH'
    archive=WORK/'transition_reports';archive.mkdir(exist_ok=True)
    for ext in ['.md','.docx']:
        if old.with_suffix(ext).exists() and not (archive/old.with_suffix(ext).name).exists():shutil.copy2(old.with_suffix(ext),archive/old.with_suffix(ext).name)
        shutil.copy2(base.with_suffix(ext),old.with_suffix(ext))
    dump(WORK/'progress_snapshot.json',{'time':time,'queue':state,'training_complete':len(trained),'smoke_complete':smoke,'inference_complete':inferred,'external_retained':220,'formal_release':release})
    print({'time':time,'stage':state['stage'],'status':state['status'],'smoke':smoke,'trained':len(trained),'active':len(state['active']),'report':str(base)})

if __name__=='__main__':main()
